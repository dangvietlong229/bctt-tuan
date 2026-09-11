import datetime as dt
import json
import re
import unittest
import tempfile
from unittest.mock import patch, Mock
from pathlib import Path

import openpyxl

import process_data
import weekly_report


ROOT = Path(__file__).resolve().parents[1]


class ReportPipelineTests(unittest.TestCase):
    def test_filename_dates_and_as_of_selection(self):
        self.assertEqual(
            weekly_report.date_from_filename(Path("source_20260821.xlsx")),
            dt.date(2026, 8, 21),
        )
        self.assertEqual(
            weekly_report.date_from_filename(Path("source_update 210826.xlsx")),
            dt.date(2026, 8, 21),
        )
        self.assertIsNone(
            weekly_report.latest_file(
                ["gd tu doanh/Thong ke Tu doanh Co phieu_update *.xlsx"],
                dt.date(2026, 8, 21),
            )
        )

    def test_missing_source_path_is_blocking(self):
        errors, _ = weekly_report.validate_sources(
            {"proprietary": ROOT / "does-not-exist.xlsx"},
            dt.date(2026, 8, 21),
            {"critical_sources": ["proprietary"], "source_max_age_days": 7},
        )
        self.assertTrue(errors)

    def test_reference_commentary_matches_required_structure(self):
        payload = json.loads(
            (ROOT / "inputs" / "commentary_reference_2026-08-24.json").read_text(encoding="utf-8")
        )
        validated = weekly_report.validate_commentary(payload)
        self.assertEqual(validated["slide8"], validated["slide3"][:2])
        self.assertTrue(validated["slide3"][0].startswith("Trong tuần "))
        self.assertTrue(validated["slide4"][0].startswith("Xét về các nhóm ngành,"))

    def test_vingroup_schema_rejects_valuation_export(self):
        files = list((ROOT / "dong gop cua vingroup").glob("FiinProX_*Doanh_nghiep*20260821*.xlsx"))
        results = {path.name: process_data.is_vingroup_price_share_export(str(path)) for path in files}
        self.assertTrue(any("(1)" in name and valid for name, valid in results.items()))
        self.assertTrue(any("(2)" in name and not valid for name, valid in results.items()))

    def test_nn_schema_rejects_valuation_export(self):
        files = list((ROOT / "nn ban rong").glob("FiinProX_*20260821*.xlsx"))
        results = {path.name: process_data.is_nn_buy_sell_room_export(str(path)) for path in files}
        self.assertTrue(any(name.endswith("20260821.xlsx") and valid for name, valid in results.items()))
        self.assertTrue(any("(2)" in name and not valid for name, valid in results.items()))

    def test_vingroup_sanity_check_rejects_corrupted_processed_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'corrupted.xlsx'
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            for row, day in enumerate((20, 21), start=3):
                sheet.cell(row, 1, dt.datetime(2026, 8, day))
                sheet.cell(row, 2, 1700)
                sheet.cell(row, 3, 7_000_000)
                for col in range(6, 14):
                    sheet.cell(row, col, 12.5)  # Valuation columns mistaken for prices/shares.
            workbook.save(path)
            workbook.close()
            self.assertIsNone(weekly_report.compute_vingroup_contribution(path, dt.date(2026, 8, 21)))

    def test_template_selection_ignores_future_and_lock_files(self):
        with tempfile.TemporaryDirectory() as directory, patch.dict('os.environ', {'REPORT_AS_OF': '2026-08-21'}):
            for name in ('Report_update 140826.xlsx', 'Report_update 280826.xlsx', '~$Report_update 210826.xlsx'):
                (Path(directory) / name).touch()
            selected = process_data.find_latest_template(directory, '*update *.xlsx', 'fallback.xlsx')
            self.assertEqual(Path(selected).name, 'Report_update 140826.xlsx')

    def test_invalid_report_date_fails(self):
        with patch.dict('os.environ', {'REPORT_AS_OF': 'bad-date'}):
            with self.assertRaises(ValueError):
                process_data.get_report_as_of()

    def test_blank_display_date(self):
        self.assertIsNone(weekly_report.parse_date('   '))
        self.assertEqual(weekly_report.display_date(''), '')

    def test_module_exclusions_reject_invalid_values(self):
        for value in ('1', '11', '2,', '2,abc'):
            with self.assertRaises(ValueError):
                process_data.parse_module_exclusions(value)
        self.assertEqual(process_data.parse_module_exclusions('2, 10,2'), {'2', '10'})

    def test_processed_date_takes_precedence(self):
        for parser in (process_data.date_from_filename, weekly_report.date_from_filename):
            self.assertEqual(parser(Path('export_20260828_update 210826.xlsx')), dt.date(2026, 8, 21))

    def test_news_empty_response_creates_header_only_workbook(self):
        response = Mock()
        response.json.return_value = {'data': []}
        with tempfile.TemporaryDirectory() as directory, patch.object(process_data.requests, 'get', return_value=response), patch.object(process_data.time, 'sleep'):
            self.assertTrue(process_data.run_tin_doanh_nghiep_feature(directory, ['ACB']))
            outputs = list(Path(directory).rglob('*.xlsx'))
            self.assertEqual(len(outputs), 1)
            workbook = openpyxl.load_workbook(outputs[0])
            self.assertEqual(workbook.active.max_row, 1)
            workbook.close()

    def test_news_failure_does_not_publish_empty_report(self):
        with tempfile.TemporaryDirectory() as directory, patch.object(process_data.requests, 'get', side_effect=RuntimeError('offline')), patch.object(process_data.time, 'sleep'):
            self.assertFalse(process_data.run_tin_doanh_nghiep_feature(directory, ['ACB']))
            self.assertFalse(list(Path(directory).rglob('*.xlsx')))

    def test_noninteractive_failure_does_not_prompt(self):
        with patch.object(process_data, 'MODULES', [('2', 'test', lambda _: False)]), patch.dict('os.environ', {'REPORT_NONINTERACTIVE': '1'}), patch('builtins.input', side_effect=AssertionError('must not prompt')):
            self.assertFalse(process_data.run_selected_modules(str(ROOT)))

    def test_vingroup_contribution_accepts_valid_price_share_rows(self):
        records = []
        for day, index, price in (
            (dt.date(2026, 8, 20), 1700.0, 100_000.0),
            (dt.date(2026, 8, 21), 1710.0, 101_000.0),
        ):
            stocks = []
            for _ in range(4):
                stocks.extend([price, 1_000_000_000.0])
            records.append({"date": day, "index": index, "market_cap": 7_000_000.0, "stocks": stocks})
        result = weekly_report.compute_vingroup_metrics(records, dt.date(2026, 8, 21))
        self.assertIsNotNone(result)
        self.assertGreater(result["week"]["group_points"], 0)
        self.assertLess(result["week"]["group_points"], 10)

    def test_output_code_has_no_calibri_or_aptos_defaults(self):
        files = [
            ROOT / "weekly_report.py",
            ROOT / "process_data.py",
            ROOT / "process_data_api.py",
            ROOT / "scripts" / "windows_office.ps1",
            ROOT / "scripts" / "build_proprietary_workbook.mjs",
        ]
        for path in files:
            text = path.read_text(encoding="utf-8")
            self.assertIsNone(
                re.search(r"(?:name\s*=\s*|name:\s*)[\"'](?:Calibri|Aptos)[\"']", text),
                path.name,
            )

    def test_font_guard_detects_old_reference_deck(self):
        path = ROOT / "output_reports" / "2026-08-21" / "MBS Dau Tu - BC Thi truong Tuan - 24.08.2026.pptx"
        with self.assertRaises(RuntimeError):
            weekly_report.assert_presentation_fonts(path)

    def test_pdf_font_guard_detects_old_reference_pdf(self):
        path = ROOT / "output_reports" / "2026-08-21" / "MBS Dau Tu - BC Thi truong Tuan - 24.08.2026.pdf"
        with self.assertRaises(RuntimeError):
            weekly_report.assert_pdf_fonts(path)

    def test_room_pdf_detection_and_flexible_date(self):
        import glob
        foreign_dir = ROOT / "gd nuoc ngoai"
        pdf_files = [
            f for f in glob.glob(str(foreign_dir / "*.pdf"))
            if Path(f).name.lower().startswith("room")
        ]
        self.assertTrue(len(pdf_files) > 0)
        as_of = dt.date(2026, 9, 4)
        def pdf_score(path):
            day = process_data.date_from_filename(path)
            diff = abs((as_of - day).days) if day else 999999
            return (diff, -Path(path).stat().st_mtime)
        selected = min(pdf_files, key=pdf_score)
        self.assertIsNotNone(selected)

    def test_run_selected_modules_skip_failed(self):
        def fail_func(ws):
            return False
        orig_modules = process_data.MODULES
        try:
            process_data.MODULES = [("99", "Mock Module", fail_func)]
            result = process_data.run_selected_modules(str(ROOT), skip_failed=True)
            self.assertTrue(result)
        finally:
            process_data.MODULES = orig_modules


if __name__ == "__main__":
    unittest.main()
