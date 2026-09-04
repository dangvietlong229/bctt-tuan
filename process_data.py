import os
import glob
import re
import datetime
import sys
import pandas as pd
import openpyxl
from openpyxl.formatting.rule import CellIsRule

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import pdfplumber
from copy import copy
import numpy as np

REPORT_FONT = "Arial"


if os.name == "nt":
    try:
        import truststore

        # Use Windows CryptoAPI roots, including organization-managed CAs.
        # SSL verification remains enabled; only the trust source changes.
        truststore.inject_into_ssl()
    except ImportError as exc:
        raise RuntimeError(
            "Missing Windows certificate support. Run 0_Cai_dat_Windows.bat, then try again."
        ) from exc


def friday_on_or_before(day: datetime.date) -> datetime.date:
    return day - datetime.timedelta(days=(day.weekday() - 4) % 7)


def get_report_as_of():
    configured = os.environ.get("REPORT_AS_OF", "").strip()
    if configured:
        try:
            return datetime.date.fromisoformat(configured)
        except ValueError:
            print(f"WARNING: REPORT_AS_OF không hợp lệ: {configured}; dùng ngày hệ thống.")
    return friday_on_or_before(datetime.date.today())


def get_update_suffix():
    return f"_update {get_report_as_of().strftime('%d%m%y')}"


def date_from_filename(file_path):
    name = os.path.basename(file_path)
    for token in re.findall(r'(?<!\d)(20\d{6})(?!\d)', name):
        try:
            return datetime.datetime.strptime(token, '%Y%m%d').date()
        except ValueError:
            pass
    for token in re.findall(r'(?i)(?:update[ _-]*)(\d{6})(?!\d)', name):
        try:
            return datetime.datetime.strptime(token, '%d%m%y').date()
        except ValueError:
            pass
    for token in re.findall(r'(?<!\d)(\d{8})(?!\d)', name):
        try:
            return datetime.datetime.strptime(token, '%d%m%Y').date()
        except ValueError:
            pass
    return None


def select_file_for_as_of(paths, as_of, require_exact=False):
    dated = [(path, date_from_filename(path)) for path in paths]
    if require_exact:
        eligible = [path for path, day in dated if day == as_of]
    else:
        eligible = [path for path, day in dated if day is not None and day <= as_of]
    if eligible:
        return max(eligible, key=lambda path: (date_from_filename(path), os.path.getmtime(path)))
    if require_exact:
        return None
    undated = [path for path, day in dated if day is None]
    return max(undated, key=os.path.getmtime) if undated else None

def save_output_file(wb, original_template_path):
    dir_name = os.path.dirname(original_template_path)
    base_name = os.path.basename(original_template_path)
    
    # Strip previous update if any to get the clean base name
    clean_name = re.sub(r'_update \d{6}', '', base_name)
    name, ext = os.path.splitext(clean_name)
    
    suffix = get_update_suffix()
    output_path = os.path.join(dir_name, f"{name}{suffix}{ext}")
    temporary_path = os.path.join(dir_name, f".{name}.{os.getpid()}.publishing{ext}")
    try:
        wb.save(temporary_path)
        os.replace(temporary_path, output_path)
    finally:
        if os.path.exists(temporary_path):
            os.remove(temporary_path)
    print(f"Saved processed output to: {output_path}")
    return output_path

def find_latest_template(directory, pattern, fallback_name):
    files = glob.glob(os.path.join(directory, pattern))
    # Exclude temporary Excel files starting with ~$
    files = [f for f in files if not os.path.basename(f).startswith('~$')]
    if files:
        # Sort by modification time to get the latest processed file
        files.sort(key=os.path.getmtime)
        return files[-1]
    return os.path.join(directory, fallback_name)

def filter_raw_files(file_list):
    return [f for f in file_list if '_update' not in os.path.basename(f)]

def translate_formulas_after_insert(ws, insert_row, num_inserted, ignore_cols=None):
    from openpyxl.formula.translate import Translator
    from openpyxl.utils import get_column_letter
    
    # Iterate through all cells starting from the first shifted row
    start_row = insert_row + num_inserted
    end_row = ws.max_row
    
    for r in range(start_row, end_row + 1):
        for c in range(1, ws.max_column + 1):
            if ignore_cols and c in ignore_cols:
                continue
            cell = ws.cell(row=r, column=c)
            val = cell.value
            if val and isinstance(val, str) and val.startswith('='):
                # Calculate the old coordinate before shifting
                old_row = r - num_inserted
                col_letter = get_column_letter(c)
                old_coord = f"{col_letter}{old_row}"
                new_coord = f"{col_letter}{r}"
                try:
                    cell.value = Translator(val, origin=old_coord).translate_formula(new_coord)
                except Exception as e:
                    print(f"Warning translating formula {val} from {old_coord} to {new_coord}: {e}")


def parse_excel_date(val):
    if not val:
        return None
    if isinstance(val, datetime.datetime):
        return val.date()
    if isinstance(val, datetime.date):
        return val
    val_str = str(val).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d"):
        try:
            return datetime.datetime.strptime(val_str, fmt).date()
        except ValueError:
            pass
    try:
        first_token = val_str.split()[0]
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d"):
            try:
                return datetime.datetime.strptime(first_token, fmt).date()
            except ValueError:
                pass
    except Exception:
        pass
    return None


def _copy_nn_cell_style(source_cell, target_cell):
    """Copy presentation only; values/formulas are populated separately."""
    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.number_format = source_cell.number_format
        target_cell.protection = copy(source_cell.protection)


def _nn_month_number(label):
    match = re.fullmatch(r'Tháng\s+(\d{1,2})', str(label or '').strip(), flags=re.IGNORECASE)
    if not match:
        return None
    month = int(match.group(1))
    return month if 1 <= month <= 12 else None


def _nn_month_summary_formula(summary_row, daily_net_column):
    date_range = "'Sheet1'!$B$10:$B$5000"
    value_range = f"'Sheet1'!${daily_net_column}$10:${daily_net_column}$5000"
    month_number = f'VALUE(SUBSTITUTE($G{summary_row},"Tháng ",""))'
    month_start = f'DATE(YEAR(MAX({date_range})),{month_number},1)'
    return (
        f'=IFERROR(SUMIFS({value_range},{date_range},">="&{month_start},'
        f'{date_range},"<"&EDATE({month_start},1)),0)'
    )


def _nn_ytd_summary_formula(daily_net_column):
    date_range = "'Sheet1'!$B$10:$B$5000"
    value_range = f"'Sheet1'!${daily_net_column}$10:${daily_net_column}$5000"
    return (
        f'=IFERROR(SUMIFS({value_range},{date_range},'
        f'">="&DATE(YEAR(MAX({date_range})),1,1),'
        f'{date_range},"<="&MAX({date_range})),0)'
    )


def update_nn_monthly_summary(ws2, latest_date, previous_latest_date=None):
    """Keep G:K aligned to month labels and formulas anchored to Sheet1 dates."""
    if not latest_date:
        return False

    summary_start_row = 5
    scan_end_row = max(ws2.max_row, 30)
    existing_month_rows = []
    ytd_row = None
    for row_idx in range(summary_start_row, scan_end_row + 1):
        label = ws2.cell(row=row_idx, column=7).value
        month = _nn_month_number(label)
        if month is not None:
            existing_month_rows.append((row_idx, month))
        elif str(label or '').strip().lower() == 'từ đầu năm':
            ytd_row = row_idx
            break

    existing_months = sorted({month for _, month in existing_month_rows})
    if previous_latest_date and previous_latest_date.year != latest_date.year:
        months_to_show = list(range(1, latest_date.month + 1))
    elif existing_months:
        first_month = min(existing_months)
        months_to_show = list(range(first_month, latest_date.month + 1))
    else:
        months_to_show = list(range(1, latest_date.month + 1))

    month_style_row = existing_month_rows[0][0] if existing_month_rows else (ytd_row or summary_start_row)
    ytd_style_row = ytd_row or month_style_row
    old_end_row = max(
        [summary_start_row - 1]
        + [row for row, _ in existing_month_rows]
        + ([ytd_row] if ytd_row else [])
    )
    new_ytd_row = summary_start_row + len(months_to_show)
    clear_end_row = max(old_end_row, new_ytd_row)

    before = [
        [ws2.cell(row=r, column=c).value for c in range(7, 12)]
        for r in range(summary_start_row, clear_end_row + 1)
    ]

    for row_idx in range(summary_start_row, clear_end_row + 1):
        for col_idx in range(7, 12):
            ws2.cell(row=row_idx, column=col_idx).value = None

    net_columns = ('E', 'J', 'O', 'T')
    for offset, month in enumerate(months_to_show):
        row_idx = summary_start_row + offset
        for col_idx in range(7, 12):
            _copy_nn_cell_style(
                ws2.cell(row=month_style_row, column=col_idx),
                ws2.cell(row=row_idx, column=col_idx),
            )
        ws2.cell(row=row_idx, column=7, value=f'Tháng {month}')
        for target_col, daily_net_col in zip(range(8, 12), net_columns):
            ws2.cell(
                row=row_idx,
                column=target_col,
                value=_nn_month_summary_formula(row_idx, daily_net_col),
            )

    for col_idx in range(7, 12):
        _copy_nn_cell_style(
            ws2.cell(row=ytd_style_row, column=col_idx),
            ws2.cell(row=new_ytd_row, column=col_idx),
        )
    ws2.cell(row=new_ytd_row, column=7, value='Từ đầu năm')
    for target_col, daily_net_col in zip(range(8, 12), net_columns):
        ws2.cell(
            row=new_ytd_row,
            column=target_col,
            value=_nn_ytd_summary_formula(daily_net_col),
        )

    after = [
        [ws2.cell(row=r, column=c).value for c in range(7, 12)]
        for r in range(summary_start_row, clear_end_row + 1)
    ]
    return before != after


def force_excel_recalculation(wb):
    """Make Excel refresh newly written formulas when the output is opened."""
    if getattr(wb, 'calculation', None) is not None:
        wb.calculation.calcMode = 'auto'
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True

# ==================== COLUMN MAPPERS ====================

def map_thanh_khoan_columns(ws_raw):
    col_map = {}
    current_group = ''
    for c in range(1, ws_raw.max_column + 1):
        group_label = str(ws_raw.cell(row=8, column=c).value or '').strip().lower()
        if group_label:
            current_group = group_label
        v9 = ws_raw.cell(row=9, column=c).value
        if not v9:
            continue
        v9_str = str(v9).lower()
        if 'ngày' in v9_str:
            col_map['date'] = c
        elif 'vnindex' in current_group:
            if 'index đóng cửa' in v9_str:
                col_map['vnindex_close'] = c
            elif 'khớp lệnh' in v9_str:
                col_map['vnindex_val'] = c
        elif 'vn30' in current_group:
            if 'index đóng cửa' in v9_str:
                col_map['vn30_close'] = c
            elif 'khớp lệnh' in v9_str:
                col_map['vn30_val'] = c
    return col_map

def map_vingroup_index_columns(ws_raw):
    col_map = {}
    current_group = ''
    for c in range(1, ws_raw.max_column + 1):
        group_label = str(ws_raw.cell(row=8, column=c).value or '').strip().lower()
        if group_label:
            current_group = group_label
        v9 = ws_raw.cell(row=9, column=c).value
        if not v9:
            continue
        v9_str = str(v9).lower()
        if 'ngày' in v9_str:
            col_map['date'] = c
        elif 'vnindex' in current_group:
            if 'index đóng cửa' in v9_str:
                col_map['vnindex_close'] = c
            elif 'vốn hóa' in v9_str:
                col_map['vnindex_mcap'] = c
    return col_map


def is_vingroup_index_export(file_path):
    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        return {'date', 'vnindex_close', 'vnindex_mcap'}.issubset(
            map_vingroup_index_columns(workbook.active)
        )
    finally:
        workbook.close()

def map_vingroup_company_columns(ws_raw):
    ticker_cols = {}
    for c in range(1, ws_raw.max_column + 1):
        t_val = ws_raw.cell(row=8, column=c).value
        field_label = str(ws_raw.cell(row=9, column=c).value or '').strip().lower()
        if t_val and isinstance(t_val, str) and 'giá đóng cửa' in field_label:
            ticker_cols[t_val.strip()] = c
    return ticker_cols


def is_vingroup_price_share_export(file_path):
    """Distinguish the required price/share export from the similarly shaped P/E-P/B export."""
    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        ticker_columns = map_vingroup_company_columns(worksheet)
        for ticker in ('VIC', 'VHM', 'VRE', 'VPL'):
            start_column = ticker_columns.get(ticker)
            if not start_column:
                return False
            price_label = str(worksheet.cell(row=9, column=start_column).value or '').strip().lower()
            shares_label = str(worksheet.cell(row=9, column=start_column + 1).value or '').strip().lower()
            if 'giá đóng cửa' not in price_label or ('cp lưu hành' not in shares_label and 'cổ phiếu lưu hành' not in shares_label):
                return False
        return True
    finally:
        workbook.close()


def write_vingroup_row(worksheet, row_index, record):
    worksheet.cell(row=row_index, column=1, value=datetime.datetime.combine(record['date'], datetime.time.min))
    worksheet.cell(row=row_index, column=2, value=record['vnindex'])
    worksheet.cell(row=row_index, column=3, value=record['mcap'])
    for ticker, price_column in (('VIC', 6), ('VHM', 8), ('VRE', 10), ('VPL', 12), ('VJC', 14)):
        worksheet.cell(row=row_index, column=price_column, value=record[ticker + '_price'])
        worksheet.cell(row=row_index, column=price_column + 1, value=record[ticker + '_shares'])

    for target_column, price_column, share_column in (
        (18, 'F', 'G'), (19, 'H', 'I'), (20, 'J', 'K'), (21, 'L', 'M'), (22, 'N', 'O')
    ):
        worksheet.cell(
            row=row_index,
            column=target_column,
            value=f"={share_column}{row_index}*{price_column}{row_index}-{share_column}{row_index+1}*{price_column}{row_index+1}",
        )
    for target_column, source_column in zip(range(24, 29), ('R', 'S', 'T', 'U', 'V')):
        worksheet.cell(row=row_index, column=target_column, value=f"={source_column}{row_index}/10^9")
    for target_column, source_column in zip(range(30, 35), ('X', 'Y', 'Z', 'AA', 'AB')):
        worksheet.cell(
            row=row_index,
            column=target_column,
            value=f"={source_column}{row_index}/$C{row_index+1}*$B{row_index+1}",
        )
    # Vingroup is VIC/VHM/VRE/VPL. VJC remains a comparison series only.
    worksheet.cell(row=row_index, column=35, value=f"=SUM(AD{row_index}:AG{row_index})")
    worksheet.cell(row=row_index, column=36, value=f"=AI{row_index}+AJ{row_index+1}")
    worksheet.cell(row=row_index, column=38, value=f"=B{row_index}-AJ{row_index}")
    worksheet.cell(row=row_index, column=39, value=f"=_xlfn.STDEV.S(AL{row_index}:AL{row_index+19})")

def map_nn_columns(ws_raw):
    ticker_cols = {}
    for c in range(1, ws_raw.max_column + 1):
        t_val = ws_raw.cell(row=8, column=c).value
        field_label = str(ws_raw.cell(row=9, column=c).value or '').strip().lower()
        if t_val and isinstance(t_val, str) and 'nđtnn mua' in field_label:
            ticker_cols[t_val.strip()] = c
    return ticker_cols


def is_nn_buy_sell_room_export(file_path):
    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        columns = map_nn_columns(worksheet)
        for ticker in ('ACB', 'VIB', 'VPB', 'TCB'):
            start = columns.get(ticker)
            if not start:
                return False
            labels = [str(worksheet.cell(row=9, column=start + offset).value or '').lower() for offset in range(3)]
            if 'nđtnn mua' not in labels[0] or 'nđtnn bán' not in labels[1] or 'còn được phép mua' not in labels[2]:
                return False
        return True
    finally:
        workbook.close()

# ==================== ACTIVE FEATURES ====================

def run_nganh_feature(workspace_dir):
    print("\n================ RUNNING NGANH FEATURE ================")
    nganh_dir = os.path.join(workspace_dir, 'nganh')
    day_files = filter_raw_files(glob.glob(os.path.join(nganh_dir, '*OneDay*.xlsx')))
    week_files = filter_raw_files(glob.glob(os.path.join(nganh_dir, '*OneWeek*.xlsx')))
    month_files = filter_raw_files(glob.glob(os.path.join(nganh_dir, '*OneMonth*.xlsx')))

    if not day_files or not week_files or not month_files:
        print("Skipping Nganh: Missing raw files for OneDay, OneWeek, or OneMonth in 'nganh' folder.")
        return False

    as_of = get_report_as_of()
    day_file = select_file_for_as_of(day_files, as_of, require_exact=True)
    week_file = select_file_for_as_of(week_files, as_of, require_exact=True)
    month_file = select_file_for_as_of(month_files, as_of, require_exact=True)
    if not day_file or not week_file or not month_file:
        print(f"Skipping Nganh: cần đủ OneDay/OneWeek/OneMonth đúng ngày {as_of:%d/%m/%Y}.")
        return False

    print(f"Reading raw files:\n- Day: {day_file}\n- Week: {week_file}\n- Month: {month_file}")

    # Helper function to load and clean industry dataframe
    def load_and_clean(file_path, value_col_name, new_col_name):
        df = pd.read_excel(file_path, header=4)
        df = df.dropna(subset=['Ngành'])
        df['Ngành'] = df['Ngành'].str.strip()
        # Filter out indexes and company information
        exclude_list = ['Tổng', 'VN30', 'VNMID', 'VNSML', 'Contact', 'CÔNG TY CỔ PHẦN FIINGROUP VIỆT NAM']
        df = df[~df['Ngành'].isin(exclude_list)]
        df_clean = df[['Ngành', value_col_name]].copy()
        df_clean.rename(columns={value_col_name: new_col_name}, inplace=True)
        return df_clean

    # Load and clean each dataset
    df_day_clean = load_and_clean(day_file, '% Thay đổi 1D', '1 ngày')
    df_week_clean = load_and_clean(week_file, '% thay đổi', '1 tuần')
    df_month_clean = load_and_clean(month_file, '% thay đổi', '1 tháng')

    # Merge datasets
    merged = pd.merge(df_day_clean, df_week_clean, on='Ngành', how='outer')
    merged = pd.merge(merged, df_month_clean, on='Ngành', how='outer')

    # Drop rows where all three columns are NaN
    merged = merged.dropna(subset=['1 ngày', '1 tuần', '1 tháng'], how='all')

    # Sort by '1 tuần' column descending
    merged = merged.sort_values(by='1 tuần', ascending=False)

    # Create new workbook for Nganh.xlsx
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'nganh'

    # Set column widths
    ws.column_dimensions['A'].width = 30
    for col in ['B', 'C', 'D']:
        ws.column_dimensions[col].width = 15

    # Merge A1:A2 and B1:D1
    ws.merge_cells('A1:A2')
    ws.merge_cells('B1:D1')

    # Set headers
    ws['A1'] = 'Ngành'
    ws['B1'] = '% Thay đổi giá'
    ws['B2'] = '1 ngày'
    ws['C2'] = '1 tuần'
    ws['D2'] = '1 tháng'

    # Stylings
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    
    # Border Side
    medium_side = Side(style='medium', color='000000')

    # Apply headers stylings
    header_coords = [
        ('A1', left_align), ('A2', left_align),
        ('B1', center_align), ('C1', center_align), ('D1', center_align),
        ('B2', center_align), ('C2', center_align), ('D2', center_align)
    ]
    for cell_ref, align in header_coords:
        cell = ws[cell_ref]
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align

    # Apply borders for header rows
    for r in [1, 2]:
        for c in range(1, 5):
            cell = ws.cell(row=r, column=c)
            left = medium_side if c == 1 else None
            right = medium_side if c == 4 else None
            top = medium_side
            bottom = medium_side
            cell.border = Border(left=left, right=right, top=top, bottom=bottom)

    # Write data rows
    data_font = Font(name=REPORT_FONT, size=11, bold=False, color='000000')
    data_left_align = Alignment(horizontal='left', vertical='center')
    data_right_align = Alignment(horizontal='right', vertical='center')

    for r_idx, row in enumerate(merged.values, start=3):
        ws.cell(row=r_idx, column=1, value=row[0])  # Ngành
        ws.cell(row=r_idx, column=2, value=row[1])  # 1 ngày
        ws.cell(row=r_idx, column=3, value=row[2])  # 1 tuần
        ws.cell(row=r_idx, column=4, value=row[3])  # 1 tháng

        # Apply formatting to data row
        for c_idx in range(1, 5):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = data_font
            
            # Alignments and number formats
            if c_idx == 1:
                cell.alignment = data_left_align
            else:
                cell.alignment = data_right_align
                cell.number_format = '0.00%'

            # Border
            left = medium_side if c_idx == 1 else None
            right = medium_side if c_idx == 4 else None
            top = medium_side
            bottom = medium_side
            cell.border = Border(left=left, right=right, top=top, bottom=bottom)

    output_path = os.path.join(workspace_dir, 'nganh', 'Nganh.xlsx')
    save_output_file(wb, output_path)
    print(f"Saved processed industry performance data to {output_path} successfully.")
    return True

def calculate_fiin_pe_pb_medians(mbs_dir):
    daily_files = glob.glob(os.path.join(mbs_dir, '*Du_lieu_giao_dich*.xlsx'))
    if not daily_files:
        daily_files = [f for f in glob.glob(os.path.join(mbs_dir, 'FiinProX_*.xlsx')) if 'du_lieu_giao_dich' in os.path.basename(f).lower()]
    if not daily_files:
        print("Warning: No daily transaction file ('*Du_lieu_giao_dich*.xlsx') found in 'danh muc mbs' for medians calculation.")
        return {}
    
    daily_file = sorted(daily_files)[-1]
    print(f"Calculating 5-year medians from daily PE/PB file: {daily_file}")
    
    try:
        wb = openpyxl.load_workbook(daily_file, read_only=True, data_only=True)
        ws = wb.active
        
        ticker_cols = {}
        col_to_metric = {}
        rows_iter = ws.iter_rows(values_only=True)
        row8 = None
        row9 = None
        for idx, row in enumerate(rows_iter, start=1):
            if idx == 8:
                row8 = row
            elif idx == 9:
                row9 = row
                break
        
        if not row8 or not row9:
            return {}
            
        curr_t = None
        for col_idx, t_val in enumerate(row8):
            if t_val and isinstance(t_val, str) and len(t_val.strip()) >= 3:
                curr_t = t_val.strip().upper()
                if curr_t not in ticker_cols:
                    ticker_cols[curr_t] = {'pe': [], 'pb': []}
            m_val = row9[col_idx] if col_idx < len(row9) else None
            if curr_t and m_val and isinstance(m_val, str):
                if 'p/e' in m_val.lower() or 'pe' in m_val.lower():
                    col_to_metric[col_idx] = (curr_t, 'pe')
                elif 'p/b' in m_val.lower() or 'pb' in m_val.lower():
                    col_to_metric[col_idx] = (curr_t, 'pb')
        
        for row in rows_iter:
            if not row or len(row) < 2 or row[1] is None:
                break
            d = row[1]
            if isinstance(d, str) and not str(d).strip().replace('-', '').isdigit():
                break
            for col_idx, (t, m_type) in col_to_metric.items():
                if col_idx < len(row):
                    v = row[col_idx]
                    if isinstance(v, (int, float)) and not np.isnan(v):
                        ticker_cols[t][m_type].append(v)
        
        wb.close()
        medians = {}
        for t, vals in ticker_cols.items():
            pe_med = float(np.median(vals['pe'])) if vals['pe'] else None
            pb_med = float(np.median(vals['pb'])) if vals['pb'] else None
            medians[t] = {'pe_median': pe_med, 'pb_median': pb_med}
        
        print(f"Calculated 5-year P/E and P/B medians for {len(medians)} stocks successfully.")
        return medians
    except Exception as e:
        print(f"Error calculating PE/PB medians from daily file: {e}")
        return {}

def run_mbs_feature(workspace_dir):
    print("\n================ RUNNING MBS FEATURE ================")
    mbs_dir = os.path.join(workspace_dir, 'danh muc mbs')
    template_path = find_latest_template(mbs_dir, 'Danh muc mbs_update *.xlsx', 'Danh muc mbs.xlsx')
    raw_files_all = filter_raw_files(glob.glob(os.path.join(mbs_dir, 'FiinProX_*.xlsx')))
    raw_files = [f for f in raw_files_all if 'du_lieu_giao_dich' not in os.path.basename(f).lower()]

    if not os.path.exists(template_path) or not raw_files:
        print("Skipping MBS: Missing template 'Danh muc mbs.xlsx' or raw data files in 'danh muc mbs' folder.")
        return False

    raw_file = select_file_for_as_of(raw_files, get_report_as_of(), require_exact=True)
    if not raw_file:
        print(f"Skipping MBS: thiếu file dữ liệu đúng ngày {get_report_as_of():%d/%m/%Y}.")
        return False
    print(f"Found template: {template_path}\nFound raw file: {raw_file}")

    medians_map = calculate_fiin_pe_pb_medians(mbs_dir)

    # Load raw file data using openpyxl to inspect headers and data
    wb_raw = openpyxl.load_workbook(raw_file, data_only=True)
    ws_raw = wb_raw.active

    # Find the header row
    header_row_idx = 8
    raw_headers = [ws_raw.cell(row=header_row_idx, column=c).value for c in range(1, ws_raw.max_column + 1)]
    
    if not any(raw_headers) or not any(x and 'mã' in str(x).lower() for x in raw_headers if isinstance(x, str)):
        for r in range(5, 12):
            row_vals = [ws_raw.cell(row=r, column=c).value for c in range(1, ws_raw.max_column + 1)]
            if any(row_vals) and any(x and 'mã' in str(x).lower() for x in row_vals if isinstance(x, str)):
                header_row_idx = r
                raw_headers = row_vals
                break

    print(f"Raw file header found at row {header_row_idx}: {raw_headers[:15]}")

    # Map headers to column indices
    col_map = {}
    for idx, val in enumerate(raw_headers):
        if val is None:
            continue
        val_clean = str(val).lower().replace('\n', ' ').strip()
        
        if 'stt' in val_clean:
            col_map['stt'] = idx + 1
        elif 'mã' in val_clean or 'ticker' in val_clean:
            col_map['ticker'] = idx + 1
        elif 'tên công ty' in val_clean or 'tên dn' in val_clean:
            col_map['company'] = idx + 1
        elif 'sàn' in val_clean:
            col_map['exchange'] = idx + 1
        elif 'giá đóng cửa' in val_clean:
            if '1 tuần' in val_clean:
                col_map['price_1w'] = idx + 1
            elif '1 tháng' in val_clean:
                col_map['price_1m'] = idx + 1
            elif '31/12' in val_clean or 'ytd' in val_clean:
                col_map['price_ytd'] = idx + 1
            else:
                col_map['price_today'] = idx + 1
        elif 'p/e' in val_clean or 'pe' in val_clean:
            col_map['pe'] = idx + 1
        elif 'p/b' in val_clean or 'pb' in val_clean:
            col_map['pb'] = idx + 1
        elif 'thấp nhất 52 tuần' in val_clean or 'low 52' in val_clean:
            col_map['low_52'] = idx + 1
        elif 'cao nhất 52 tuần' in val_clean or 'high 52' in val_clean:
            col_map['high_52'] = idx + 1

    print("Mapped columns:", col_map)

    required_cols = {'stt', 'ticker', 'price_today', 'price_1w'}
    if not required_cols.issubset(col_map.keys()):
        print(f"Skipping MBS: Missing required columns in raw file. Found mapping: {col_map.keys()}")
        wb_raw.close()
        return False

    raw_data_list = []
    for r in range(header_row_idx + 1, ws_raw.max_row + 1):
        ticker = ws_raw.cell(row=r, column=col_map['ticker']).value
        if ticker and isinstance(ticker, str):
            ticker = ticker.strip().upper()
            if len(ticker) < 3 or len(ticker) > 10:
                continue
            
            def get_val(col_name):
                col_idx = col_map.get(col_name)
                return ws_raw.cell(row=r, column=col_idx).value if col_idx else None

            raw_data_list.append({
                'stt': get_val('stt'),
                'ticker': ticker,
                'company': get_val('company'),
                'exchange': get_val('exchange'),
                'price_today': get_val('price_today'),
                'price_1w': get_val('price_1w'),
                'price_1m': get_val('price_1m'),
                'price_ytd': get_val('price_ytd'),
                'pe': get_val('pe'),
                'pb': get_val('pb'),
                'low_52': get_val('low_52'),
                'high_52': get_val('high_52'),
            })

    wb_raw.close()
    print(f"Read {len(raw_data_list)} valid data rows from raw file.")

    wb_dest = openpyxl.load_workbook(template_path)
    if "Cập nhật giá" not in wb_dest.sheetnames or "Danh mục" not in wb_dest.sheetnames:
        print("Skipping MBS: Sheets 'Cập nhật giá' or 'Danh mục' not found in template.")
        wb_dest.close()
        return False

    ws_price = wb_dest["Cập nhật giá"]

    # Copy style from row 9 (before deleting)
    styles = {}
    for col_idx in range(1, 18):
        cell = ws_price.cell(row=9, column=col_idx)
        styles[col_idx] = {
            'font': cell.font,
            'fill': cell.fill,
            'border': cell.border,
            'alignment': cell.alignment,
            'number_format': cell.number_format
        }
        if col_idx in (16, 17) and (not styles[col_idx]['font'] or not styles[col_idx]['font'].name):
            base_col = 12 if col_idx == 16 else 13
            if base_col in styles:
                styles[col_idx] = copy(styles[base_col])

    # Delete existing data from row 9 to the end
    if ws_price.max_row >= 9:
        ws_price.delete_rows(9, ws_price.max_row - 9 + 1)

    ws_price.cell(row=8, column=16, value="P/E Trung vị 5 năm\nĐơn vị: Lần")
    ws_price.cell(row=8, column=17, value="P/B Trung vị 5 năm\nĐơn vị: Lần")
    ws_price.cell(row=8, column=14, value="Giá cao nhất 52 tuần\nĐơn vị: VND")
    ws_price.cell(row=8, column=15, value="Giá thấp nhất 52 tuần\nĐơn vị: VND")

    # Write new data and reapply formulas & styles
    for idx, data_row in enumerate(raw_data_list):
        row_idx = 9 + idx
        ws_price.cell(row=row_idx, column=1, value=data_row['stt'])
        ws_price.cell(row=row_idx, column=2, value=data_row['ticker'])
        ws_price.cell(row=row_idx, column=3, value=data_row['company'])
        ws_price.cell(row=row_idx, column=4, value=data_row['exchange'])
        ws_price.cell(row=row_idx, column=5, value=data_row['price_today'])
        ws_price.cell(row=row_idx, column=6, value=data_row['price_1w'])
        ws_price.cell(row=row_idx, column=7, value=data_row['price_1m'])
        ws_price.cell(row=row_idx, column=8, value=data_row['price_ytd'])
        
        # Formulas
        ws_price.cell(row=row_idx, column=9, value=f"=$E{row_idx}/F{row_idx}-1")
        ws_price.cell(row=row_idx, column=10, value=f"=$E{row_idx}/G{row_idx}-1")
        ws_price.cell(row=row_idx, column=11, value=f"=$E{row_idx}/H{row_idx}-1")
        
        ws_price.cell(row=row_idx, column=12, value=data_row['pe'])
        ws_price.cell(row=row_idx, column=13, value=data_row['pb'])
        # "Danh mục" expects 52w high first (N), then 52w low (O).
        ws_price.cell(row=row_idx, column=14, value=data_row['high_52'])
        ws_price.cell(row=row_idx, column=15, value=data_row['low_52'])

        pe_med = medians_map.get(data_row['ticker'], {}).get('pe_median')
        pb_med = medians_map.get(data_row['ticker'], {}).get('pb_median')
        ws_price.cell(row=row_idx, column=16, value=pe_med)
        ws_price.cell(row=row_idx, column=17, value=pb_med)

        # Apply copied styles
        for col_idx in range(1, 18):
            cell = ws_price.cell(row=row_idx, column=col_idx)
            if col_idx in styles:
                s = styles[col_idx]
                if s['font']: cell.font = copy(s['font'])
                if s['fill']: cell.fill = copy(s['fill'])
                if s['border']: cell.border = copy(s['border'])
                if s['alignment']: cell.alignment = copy(s['alignment'])
                if s['number_format']: cell.number_format = s['number_format']

    print(f"Updated {len(raw_data_list)} rows in 'Cập nhật giá' sheet.")

    # Sort sheet "Danh mục"
    ws_danh_muc = wb_dest["Danh mục"]
    
    danh_muc_rows = []
    for r in range(4, ws_danh_muc.max_row + 1):
        ticker = ws_danh_muc.cell(row=r, column=2).value
        if ticker:
            stt = ws_danh_muc.cell(row=r, column=1).value
            danh_muc_rows.append({
                'ticker': str(ticker).strip().upper(),
                'stt': stt
            })

    weekly_return_map = {}
    for d in raw_data_list:
        ticker = d['ticker']
        today = d['price_today']
        w1 = d['price_1w']
        if today is not None and w1 is not None and w1 != 0:
            try:
                weekly_return_map[ticker] = float(today) / float(w1) - 1.0
            except (ValueError, TypeError):
                weekly_return_map[ticker] = -999.0
        else:
            weekly_return_map[ticker] = -999.0

    for item in danh_muc_rows:
        item['return'] = weekly_return_map.get(item['ticker'], -999.0)

    danh_muc_rows.sort(key=lambda x: x['return'], reverse=True)

    print(f"Sorting 'Danh mục' sheet. Top 5 tickers after sorting:")
    for item in danh_muc_rows[:5]:
        print(f"  Ticker: {item['ticker']}, return: {item['return']:.4f}")

    lookup_end_row = max(9, 8 + len(raw_data_list))
    lookup_range = f"'Cập nhật giá'!$B$9:$Q${lookup_end_row}"
    lookup_columns = {
        3: 4,
        4: 8,
        5: 9,
        6: 10,
        9: 11,
        10: 12,
        13: 13,
        14: 14,
        15: 15,
        16: 16,
    }

    for idx, item in enumerate(danh_muc_rows):
        r = 4 + idx
        # The list is re-sorted every run, so the source STT is no longer a valid
        # display order.  Always number the 70-row output sequentially.
        ws_danh_muc.cell(row=r, column=1, value=idx + 1)
        ws_danh_muc.cell(row=r, column=2, value=item['ticker'])

        for target_col, source_index in lookup_columns.items():
            ws_danh_muc.cell(
                row=r,
                column=target_col,
                value=f"=VLOOKUP($B{r},{lookup_range},{source_index},FALSE)",
            )
        ws_danh_muc.cell(row=r, column=7, value=f"=C{r}/M{r}-1")
        ws_danh_muc.cell(row=r, column=8, value=f"=C{r}/N{r}-1")
        ws_danh_muc.cell(row=r, column=11, value=f"=I{r}/O{r}-1")
        ws_danh_muc.cell(row=r, column=12, value=f"=J{r}/P{r}-1")

    # Rebuild the percentage conditional formatting on every run.  Formula cells
    # keep their normal style; Excel evaluates these rules after recalculation.
    for cell_range in ("D4:H73", "K4:L73"):
        for existing in list(ws_danh_muc.conditional_formatting._cf_rules):
            if str(existing.sqref) == cell_range:
                del ws_danh_muc.conditional_formatting._cf_rules[existing]
        ws_danh_muc.conditional_formatting.add(
            cell_range, CellIsRule(operator="lessThan", formula=["0"], font=Font(color="9C0006"))
        )
        ws_danh_muc.conditional_formatting.add(
            cell_range, CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="00B050"))
        )

    save_output_file(wb_dest, template_path)
    print(f"Updated and saved 'Danh mục' sheets in {template_path} successfully.")
    return True

def run_foreign_feature(workspace_dir):
    print("\n================ RUNNING GD NUOC NGOAI FEATURE ================")
    foreign_dir = os.path.join(workspace_dir, 'gd nuoc ngoai')
    pdf_files = glob.glob(os.path.join(foreign_dir, 'Room_*.pdf'))
    ban_files = filter_raw_files(glob.glob(os.path.join(foreign_dir, 'FiinProX_*Ban*.xlsx')))
    mua_files = filter_raw_files(glob.glob(os.path.join(foreign_dir, 'FiinProX_*Mua*.xlsx')))

    as_of = get_report_as_of()
    ban_file = select_file_for_as_of(ban_files, as_of, require_exact=True)
    mua_file = select_file_for_as_of(mua_files, as_of, require_exact=True)
    pdf_file = select_file_for_as_of(pdf_files, as_of, require_exact=False)

    if not pdf_file or not ban_file or not mua_file:
        print(
            "Skipping GD Nuoc Ngoai: cần đủ Room PDF và cặp Mua/Bán đúng ngày "
            f"{as_of:%d/%m/%Y}."
        )
        return False
    if date_from_filename(ban_file) != date_from_filename(mua_file):
        print("Skipping GD Nuoc Ngoai: file Mua và Bán không cùng ngày dữ liệu.")
        return False
    room_date = date_from_filename(pdf_file)
    if room_date and (as_of - room_date).days > 31:
        print(f"Skipping GD Nuoc Ngoai: Room PDF đã cũ {(as_of - room_date).days} ngày.")
        return False
    if room_date and room_date != as_of:
        print(f"WARNING: dùng Room PDF ngày {room_date:%d/%m/%Y} cho kỳ {as_of:%d/%m/%Y}.")
    print(f"Found room PDF: {pdf_file}")

    room_data = {}
    print("Parsing PDF room file...")
    with pdfplumber.open(pdf_file) as pdf:
        for idx, page in enumerate(pdf.pages):
            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    if len(row) >= 8:
                        ticker = row[1]
                        if ticker and isinstance(ticker, str):
                            ticker = ticker.strip()
                            if ticker.isupper() and 3 <= len(ticker) <= 8:
                                total_issued_str = row[3]
                                owned_qty_str = row[6]
                                if total_issued_str and owned_qty_str:
                                    try:
                                        total_issued = float(total_issued_str.replace('.', '').replace(',', ''))
                                        owned_qty = float(owned_qty_str.replace('.', '').replace(',', ''))
                                        if total_issued > 0:
                                            pct = owned_qty / total_issued
                                            room_data[ticker] = pct
                                    except ValueError:
                                        pass

    print(f"Successfully extracted room data for {len(room_data)} tickers from PDF.")

    def update_excel_file(excel_path):
        print(f"Updating file: {excel_path}")
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active

        header_row = 5
        ticker_col = 1
        pct_col = 7
        
        found_ticker = False
        found_pct = False
        for r in range(1, 10):
            for c in range(1, ws.max_column + 1):
                val = ws.cell(row=r, column=c).value
                if val:
                    val_str = str(val).lower().strip()
                    if val_str == 'mã' or val_str == 'ticker':
                        header_row = r
                        ticker_col = c
                        found_ticker = True
                    elif 'sở hữu' in val_str:
                        pct_col = c
                        found_pct = True
            if found_ticker and found_pct:
                break

        print(f"  Header row: {header_row}, Ticker column: {ticker_col}, % Sở hữu column: {pct_col}")

        updated_count = 0
        for r in range(header_row + 1, ws.max_row + 1):
            ticker = ws.cell(row=r, column=ticker_col).value
            if ticker and isinstance(ticker, str):
                ticker = ticker.strip()
                if ticker in room_data:
                    pct_val = room_data[ticker]
                    cell = ws.cell(row=r, column=pct_col, value=pct_val)
                    cell.number_format = '0.00%'
                    updated_count += 1
                else:
                    if len(ticker) == 3:
                        print(f"  Warning: Stock Ticker {ticker} not found in PDF room data.")
        
        dir_name = os.path.dirname(excel_path)
        base_name = os.path.basename(excel_path)
        clean_name = re.sub(r'_update \d{6}', '', base_name)
        name, ext = os.path.splitext(clean_name)
        suffix = get_update_suffix()
        suffix_path = os.path.join(dir_name, f"{name}{suffix}{ext}")
        temporary_path = os.path.join(dir_name, f".{name}.{os.getpid()}.publishing{ext}")
        try:
            wb.save(temporary_path)
            os.replace(temporary_path, suffix_path)
        finally:
            wb.close()
            if os.path.exists(temporary_path):
                os.remove(temporary_path)
        print(f"  Successfully updated {updated_count} rows and saved to {os.path.basename(suffix_path)}")
        return suffix_path

    new_saved_paths = []
    for f in [ban_file]:
        p = update_excel_file(f)
        if p:
            new_saved_paths.append(os.path.abspath(p))

    for f in [mua_file]:
        p = update_excel_file(f)
        if p:
            new_saved_paths.append(os.path.abspath(p))

    return True

def run_thanh_khoan_feature(workspace_dir):
    print("\n================ RUNNING THANH KHOAN TT FEATURE ================")
    tk_dir = os.path.join(workspace_dir, 'thanh khoan tt')
    template_path = find_latest_template(tk_dir, 'Thanh khoan tt_update *.xlsx', 'Thanh khoan tt.xlsx')
    raw_files = filter_raw_files(glob.glob(os.path.join(tk_dir, 'FiinProX_*.xlsx')))

    if not os.path.exists(template_path) or not raw_files:
        print("Skipping Thanh Khoan: Missing template or raw data files.")
        return False

    raw_file = select_file_for_as_of(raw_files, get_report_as_of(), require_exact=True)
    if not raw_file:
        print(f"Skipping Thanh Khoan: thiếu file dữ liệu đúng ngày {get_report_as_of():%d/%m/%Y}.")
        return False
    print(f"Found template: {template_path}\nFound raw file: {raw_file}")

    wb_raw = openpyxl.load_workbook(raw_file, data_only=True)
    ws_raw = wb_raw.active
    col_map = map_thanh_khoan_columns(ws_raw)

    required_keys = {'date', 'vnindex_close', 'vnindex_val', 'vn30_close', 'vn30_val'}
    if not required_keys.issubset(col_map.keys()):
        print(f"Skipping Thanh Khoan: Missing required columns in raw file. Found: {col_map.keys()}")
        return False

    raw_data = []
    for r in range(10, ws_raw.max_row + 1):
        date_val = ws_raw.cell(row=r, column=col_map['date']).value
        if date_val:
            if isinstance(date_val, str) and ('contact' in date_val.lower() or 'công ty' in date_val.lower()):
                continue
            
            # Parse date safely
            dt = parse_excel_date(date_val)
            if not dt:
                continue

            raw_data.append({
                'date': dt,
                'vnindex_close': ws_raw.cell(row=r, column=col_map['vnindex_close']).value,
                'vnindex_val': ws_raw.cell(row=r, column=col_map['vnindex_val']).value,
                'vn30_close': ws_raw.cell(row=r, column=col_map['vn30_close']).value,
                'vn30_val': ws_raw.cell(row=r, column=col_map['vn30_val']).value
            })

    wb_dest = openpyxl.load_workbook(template_path, data_only=False)
    ws = wb_dest.active

    # Find max date currently in template
    template_dates = []
    for r in range(10, ws.max_row + 1):
        val = ws.cell(row=r, column=2).value
        if val:
            dt = parse_excel_date(val)
            if dt:
                template_dates.append(dt)

    max_template_date = max(template_dates) if template_dates else None
    print(f"Max template date: {max_template_date}")

    # Filter to only insert new dates
    new_rows = []
    for row in raw_data:
        if max_template_date is None or row['date'] > max_template_date:
            new_rows.append(row)

    N = len(new_rows)
    print(f"New rows to insert: {N}")

    if N > 0:
        new_rows.sort(key=lambda x: x['date'], reverse=True)
        
        # Copy formatting of row 10
        styles = {}
        for col_idx in range(1, 13):
            cell = ws.cell(row=10, column=col_idx)
            styles[col_idx] = {
                'font': cell.font,
                'fill': cell.fill,
                'border': cell.border,
                'alignment': cell.alignment,
                'number_format': cell.number_format
            }

        ws.insert_rows(10, N)
        translate_formulas_after_insert(ws, 10, N)

        # Write data
        for idx, row in enumerate(new_rows):
            r_idx = 10 + idx
            ws.cell(row=r_idx, column=1, value='')
            ws.cell(row=r_idx, column=2, value=datetime.datetime.combine(row['date'], datetime.time.min))
            ws.cell(row=r_idx, column=3, value=row['vnindex_close'])
            ws.cell(row=r_idx, column=4, value=row['vnindex_val'])
            ws.cell(row=r_idx, column=5, value=row['vn30_close'])
            ws.cell(row=r_idx, column=6, value=row['vn30_val'])

            # Write formulas
            wday = row['date'].weekday()
            if wday == 4:  # Friday
                ws.cell(row=r_idx, column=8, value=f"=C{r_idx}-C{r_idx+5}")
                ws.cell(row=r_idx, column=9, value=f"=C{r_idx}/C{r_idx+5}-1")
                ws.cell(row=r_idx, column=10, value=f"=E{r_idx}-E{r_idx+5}")
                ws.cell(row=r_idx, column=11, value=f"=E{r_idx}/E{r_idx+5}-1")
            elif wday == 0:  # Monday
                ws.cell(row=r_idx, column=8, value=f"=SUM(D{r_idx-4}:D{r_idx})")
                ws.cell(row=r_idx, column=9, value=f"=H{r_idx}/H{r_idx+5}-1")
                ws.cell(row=r_idx, column=10, value=f"=SUM(F{r_idx-4}:F{r_idx})")
                ws.cell(row=r_idx, column=11, value=f"=J{r_idx}/J{r_idx+5}-1")

            # Apply style
            for col_idx in range(1, 13):
                cell = ws.cell(row=r_idx, column=col_idx)
                if col_idx in styles:
                    s = styles[col_idx]
                    if s['font']: cell.font = copy(s['font'])
                    if s['fill']: cell.fill = copy(s['fill'])
                    if s['border']: cell.border = copy(s['border'])
                    if s['alignment']: cell.alignment = copy(s['alignment'])
                    if s['number_format']: cell.number_format = s['number_format']

        # Re-number STT from row 15 downwards
        stt = 1
        for r in range(15, ws.max_row + 1):
            date_val = ws.cell(row=r, column=2).value
            if date_val:
                ws.cell(row=r, column=1, value=stt)
                stt += 1
            else:
                break

        save_output_file(wb_dest, template_path)
        print(f"Successfully processed Thanh Khoan with {N} new rows.")
    else:
        print("No new dates to insert for Thanh Khoan.")
    return True

def run_vingroup_feature(workspace_dir):
    print("\n================ RUNNING DONG GOP CUA VINGROUP FEATURE ================")
    vg_dir = os.path.join(workspace_dir, 'dong gop cua vingroup')
    template_path = find_latest_template(vg_dir, 'Dong gop cua vingroup_update *.xlsx', 'Dong gop cua vingroup.xlsx')
    index_files = filter_raw_files(glob.glob(os.path.join(vg_dir, 'FiinProX_*Chi_so_*.xlsx')))
    comp_files = filter_raw_files(glob.glob(os.path.join(vg_dir, 'FiinProX_*Doanh_nghiep*.xlsx')))

    if not os.path.exists(template_path) or not index_files or not comp_files:
        print("Skipping Vingroup: Missing template or raw index/company files.")
        return False

    as_of = get_report_as_of()
    valid_index_files = [path for path in index_files if is_vingroup_index_export(path)]
    index_file = select_file_for_as_of(valid_index_files, as_of, require_exact=True)
    valid_company_files = [path for path in comp_files if is_vingroup_price_share_export(path)]
    comp_file = select_file_for_as_of(valid_company_files, as_of, require_exact=True)
    if not index_file or not comp_file:
        print("Skipping Vingroup: không tìm thấy đúng file chỉ số và file Giá đóng cửa/Số CP lưu hành.")
        return False
    print(f"Index raw file: {index_file}\nCompany raw file: {comp_file}")

    # Parse index raw file
    wb_idx = openpyxl.load_workbook(index_file, data_only=True)
    ws_idx = wb_idx.active
    idx_map = map_vingroup_index_columns(ws_idx)

    required_index_keys = {'date', 'vnindex_close', 'vnindex_mcap'}
    missing_index_keys = sorted(required_index_keys - set(idx_map))
    if missing_index_keys:
        labels = {
            'date': 'Ngày',
            'vnindex_close': 'Index đóng cửa của VNINDEX',
            'vnindex_mcap': 'Vốn hóa thị trường của VNINDEX',
        }
        missing_labels = ', '.join(labels[key] for key in missing_index_keys)
        print(
            "Skipping Vingroup: File dữ liệu chỉ số thiếu cột bắt buộc: "
            f"{missing_labels}. Hãy xuất lại dữ liệu FiinProX có trường Vốn hóa thị trường."
        )
        wb_idx.close()
        return False

    index_data = {}
    for r in range(10, ws_idx.max_row + 1):
        d_val = ws_idx.cell(row=r, column=idx_map['date']).value
        if d_val:
            if isinstance(d_val, str) and ('contact' in d_val.lower() or 'công ty' in d_val.lower()):
                continue
            dt = parse_excel_date(d_val)
            if not dt:
                continue

            close_val = ws_idx.cell(row=r, column=idx_map['vnindex_close']).value
            mcap_val = ws_idx.cell(row=r, column=idx_map['vnindex_mcap']).value

            try:
                mcap_bil = round(float(mcap_val) / 1e9)
            except (ValueError, TypeError):
                mcap_bil = None

            index_data[dt] = {
                'vnindex': close_val,
                'mcap': mcap_bil
            }
    wb_idx.close()

    # Parse company raw file
    wb_comp = openpyxl.load_workbook(comp_file, data_only=True)
    ws_comp = wb_comp.active
    comp_map = map_vingroup_company_columns(ws_comp)

    c_date_col = 2
    company_data = {}
    for r in range(10, ws_comp.max_row + 1):
        d_val = ws_comp.cell(row=r, column=c_date_col).value
        if d_val:
            if isinstance(d_val, str) and ('contact' in d_val.lower() or 'công ty' in d_val.lower()):
                continue
            dt = parse_excel_date(d_val)
            if not dt:
                continue

            row_info = {}
            for ticker in ['VIC', 'VHM', 'VRE', 'VPL', 'VJC']:
                start_col = comp_map.get(ticker)
                if start_col:
                    row_info[ticker + '_price'] = ws_comp.cell(row=r, column=start_col).value
                    row_info[ticker + '_shares'] = ws_comp.cell(row=r, column=start_col + 1).value
                else:
                    row_info[ticker + '_price'] = None
                    row_info[ticker + '_shares'] = None
            company_data[dt] = row_info
    wb_comp.close()

    # Merge on Date
    merged_raw = []
    for dt in sorted(index_data.keys()):
        if dt in company_data:
            merged_raw.append({
                'date': dt,
                'vnindex': index_data[dt]['vnindex'],
                'mcap': index_data[dt]['mcap'],
                **company_data[dt]
            })

    merged_raw.sort(key=lambda x: x['date'], reverse=True)
    print(f"Merged raw records: {len(merged_raw)}")

    wb_dest = openpyxl.load_workbook(template_path, data_only=False)
    ws = wb_dest.active

    # Find max date in template
    template_dates = []
    for r in range(3, ws.max_row + 1):
        val = ws.cell(row=r, column=1).value
        if val:
            dt = parse_excel_date(val)
            if dt:
                template_dates.append(dt)

    max_template_date = max(template_dates) if template_dates else None
    print(f"Max template date: {max_template_date}")

    new_rows = [row for row in merged_raw if max_template_date is None or row['date'] > max_template_date]
    N = len(new_rows)
    print(f"New rows to insert: {N}")

    if N > 0:
        # Copy style from row 3
        styles = {}
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=3, column=col_idx)
            styles[col_idx] = {
                'font': cell.font,
                'fill': cell.fill,
                'border': cell.border,
                'alignment': cell.alignment,
                'number_format': cell.number_format
            }

        ws.insert_rows(3, N)
        translate_formulas_after_insert(ws, 3, N, ignore_cols=range(40, 47))

        # Write data and formulas
        for idx, row in enumerate(new_rows):
            r_idx = 3 + idx
            write_vingroup_row(ws, r_idx, row)

            # Apply style
            for col_idx in range(1, ws.max_column + 1):
                if 40 <= col_idx <= 46:
                    continue
                cell = ws.cell(row=r_idx, column=col_idx)
                if col_idx in styles:
                    s = styles[col_idx]
                    if s['font']: cell.font = copy(s['font'])
                    if s['fill']: cell.fill = copy(s['fill'])
                    if s['border']: cell.border = copy(s['border'])
                    if s['alignment']: cell.alignment = copy(s['alignment'])
                    if s['number_format']: cell.number_format = s['number_format']

        print(f"Inserted {N} new Vingroup rows.")
    else:
        print("No new dates to insert for Vingroup.")

    # Repair all overlapping dates as well as adding new ones. This prevents a
    # previously selected P/E-P/B export from permanently contaminating history.
    date_rows = {}
    for row_index in range(3, ws.max_row + 1):
        day = parse_excel_date(ws.cell(row=row_index, column=1).value)
        if day:
            date_rows[day] = row_index
    repaired = 0
    for record in merged_raw:
        row_index = date_rows.get(record['date'])
        if row_index is not None:
            write_vingroup_row(ws, row_index, record)
            repaired += 1
    force_excel_recalculation(wb_dest)
    save_output_file(wb_dest, template_path)
    wb_dest.close()
    print(f"Successfully upserted {repaired} Vingroup rows from the validated price/share export.")
    return True

def run_nn_ban_rong_feature(workspace_dir):
    print("\n================ RUNNING NN BAN RONG FEATURE ================")
    nn_dir = os.path.join(workspace_dir, 'nn ban rong')
    template_path = find_latest_template(nn_dir, 'NN ban rong_update *.xlsx', 'NN ban rong.xlsx')
    raw_files = filter_raw_files(glob.glob(os.path.join(nn_dir, 'FiinProX_*.xlsx')))

    if not os.path.exists(template_path) or not raw_files:
        print("Skipping NN ban rong: Missing template or raw data files.")
        return False

    valid_raw_files = [path for path in raw_files if is_nn_buy_sell_room_export(path)]
    raw_file = select_file_for_as_of(valid_raw_files, get_report_as_of(), require_exact=True)
    if not raw_file:
        print(f"Skipping NN ban rong: thiếu file Mua/Bán/Room đúng ngày {get_report_as_of():%d/%m/%Y}.")
        return False
    print(f"Found template: {template_path}\nFound raw file: {raw_file}")

    wb_raw = openpyxl.load_workbook(raw_file, data_only=True)
    ws_raw = wb_raw.active
    comp_map = map_nn_columns(ws_raw)

    raw_data = []
    for r in range(10, ws_raw.max_row + 1):
        date_val = ws_raw.cell(row=r, column=2).value
        if date_val:
            if isinstance(date_val, str) and ('contact' in date_val.lower() or 'công ty' in date_val.lower()):
                continue
            dt = parse_excel_date(date_val)
            if not dt:
                continue

            row_info = {'date': dt}
            for ticker in ['ACB', 'VIB', 'VPB', 'TCB']:
                start_col = comp_map.get(ticker)
                if start_col:
                    row_info[ticker + '_buy'] = ws_raw.cell(row=r, column=start_col).value
                    row_info[ticker + '_sell'] = ws_raw.cell(row=r, column=start_col + 1).value
                    row_info[ticker + '_room'] = ws_raw.cell(row=r, column=start_col + 2).value
                else:
                    row_info[ticker + '_buy'] = None
                    row_info[ticker + '_sell'] = None
                    row_info[ticker + '_room'] = None
            raw_data.append(row_info)

    raw_data.sort(key=lambda x: x['date'], reverse=True)

    # Load template with data_only=True to get the last cumulative values
    wb_dest_val = openpyxl.load_workbook(template_path, data_only=True)
    ws1_val = wb_dest_val["Sheet1"]
    ws2_val = wb_dest_val["Sheet2"]

    # Find max date in template Sheet1
    template_dates = []
    for r in range(10, ws1_val.max_row + 1):
        val = ws1_val.cell(row=r, column=2).value
        if val:
            dt = parse_excel_date(val)
            if dt:
                template_dates.append(dt)
    max_template_date = max(template_dates) if template_dates else None
    print(f"Max template date: {max_template_date}")

    cum_acb = 0.0
    cum_vib = 0.0
    cum_vpb = 0.0
    cum_tcb = 0.0

    found_row = None
    if max_template_date:
        for r in range(5, ws2_val.max_row + 1):
            val = ws2_val.cell(row=r, column=1).value
            if val:
                dt_val = parse_excel_date(val)
                if dt_val == max_template_date:
                    found_row = r
                    break

    if found_row:
        try:
            cum_acb = float(ws2_val.cell(row=found_row, column=2).value or 0.0)
            cum_vib = float(ws2_val.cell(row=found_row, column=3).value or 0.0)
            cum_vpb = float(ws2_val.cell(row=found_row, column=4).value or 0.0)
            cum_tcb = float(ws2_val.cell(row=found_row, column=5).value or 0.0)
            print(f"Found starting cumulative values in Sheet2 row {found_row} for date {max_template_date}: ACB={cum_acb}, VIB={cum_vib}, VPB={cum_vpb}, TCB={cum_tcb}")
        except (ValueError, TypeError) as e:
            print(f"Error parsing cumulative values from Sheet2 row {found_row}: {e}. Defaulting to 0.0")
    else:
        print(f"Could not find cumulative values in Sheet2 matching max template date {max_template_date}. Defaulting to 0.0")

    wb_dest = openpyxl.load_workbook(template_path, data_only=False)
    ws1 = wb_dest["Sheet1"]
    ws2 = wb_dest["Sheet2"]

    new_rows = [row for row in raw_data if max_template_date is None or row['date'] > max_template_date]
    N = len(new_rows)
    print(f"New rows to insert: {N}")

    if N > 0:
        # Calculate cumulative values
        new_rows_asc = sorted(new_rows, key=lambda x: x['date'])
        cum_values = {}
        for row in new_rows_asc:
            dt = row['date']
            acb_net = (float(row['ACB_buy']) if row['ACB_buy'] is not None else 0.0) - (float(row['ACB_sell']) if row['ACB_sell'] is not None else 0.0)
            vib_net = (float(row['VIB_buy']) if row['VIB_buy'] is not None else 0.0) - (float(row['VIB_sell']) if row['VIB_sell'] is not None else 0.0)
            vpb_net = (float(row['VPB_buy']) if row['VPB_buy'] is not None else 0.0) - (float(row['VPB_sell']) if row['VPB_sell'] is not None else 0.0)
            tcb_net = (float(row['TCB_buy']) if row['TCB_buy'] is not None else 0.0) - (float(row['TCB_sell']) if row['TCB_sell'] is not None else 0.0)
            
            cum_acb += acb_net
            cum_vib += vib_net
            cum_vpb += vpb_net
            cum_tcb += tcb_net
            
            cum_values[dt] = {
                'ACB': cum_acb,
                'VIB': cum_vib,
                'VPB': cum_vpb,
                'TCB': cum_tcb
            }

        # 3. Update Sheet1
        # Copy style of row 10 in Sheet1
        styles1 = {}
        for col_idx in range(1, 23):
            cell = ws1.cell(row=10, column=col_idx)
            styles1[col_idx] = {
                'font': cell.font,
                'fill': cell.fill,
                'border': cell.border,
                'alignment': cell.alignment,
                'number_format': cell.number_format
            }

        ws1.insert_rows(10, N)
        translate_formulas_after_insert(ws1, 10, N)

        # Write data to Sheet1 (newest first)
        for idx, row in enumerate(new_rows):
            r_idx = 10 + idx
            ws1.cell(row=r_idx, column=1, value='')
            ws1.cell(row=r_idx, column=2, value=datetime.datetime.combine(row['date'], datetime.time.min))
            
            # ACB
            ws1.cell(row=r_idx, column=3, value=row['ACB_buy'])
            ws1.cell(row=r_idx, column=4, value=row['ACB_sell'])
            ws1.cell(row=r_idx, column=5, value=f"=C{r_idx}-D{r_idx}")
            ws1.cell(row=r_idx, column=6, value=f"=F{r_idx+1}+E{r_idx}")
            ws1.cell(row=r_idx, column=7, value=row['ACB_room'])
            
            # VIB
            ws1.cell(row=r_idx, column=8, value=row['VIB_buy'])
            ws1.cell(row=r_idx, column=9, value=row['VIB_sell'])
            ws1.cell(row=r_idx, column=10, value=f"=H{r_idx}-I{r_idx}")
            ws1.cell(row=r_idx, column=11, value=f"=K{r_idx+1}+J{r_idx}")
            ws1.cell(row=r_idx, column=12, value=row['VIB_room'])
            
            # VPB
            ws1.cell(row=r_idx, column=13, value=row['VPB_buy'])
            ws1.cell(row=r_idx, column=14, value=row['VPB_sell'])
            ws1.cell(row=r_idx, column=15, value=f"=M{r_idx}-N{r_idx}")
            ws1.cell(row=r_idx, column=16, value=f"=P{r_idx+1}+O{r_idx}")
            ws1.cell(row=r_idx, column=17, value=row['VPB_room'])
            
            # TCB
            ws1.cell(row=r_idx, column=18, value=row['TCB_buy'])
            ws1.cell(row=r_idx, column=19, value=row['TCB_sell'])
            ws1.cell(row=r_idx, column=20, value=f"=R{r_idx}-S{r_idx}")
            ws1.cell(row=r_idx, column=21, value=f"=U{r_idx+1}+T{r_idx}")
            ws1.cell(row=r_idx, column=22, value=row['TCB_room'])

            for col_idx in range(1, 23):
                cell = ws1.cell(row=r_idx, column=col_idx)
                if col_idx in styles1:
                    s = styles1[col_idx]
                    if s['font']: cell.font = copy(s['font'])
                    if s['fill']: cell.fill = copy(s['fill'])
                    if s['border']: cell.border = copy(s['border'])
                    if s['alignment']: cell.alignment = copy(s['alignment'])
                    if s['number_format']: cell.number_format = s['number_format']

        # 4. Update Sheet2
        # Style for column 1 to 5 of row 5 (to copy style to new rows)
        styles2 = {}
        for col_idx in range(1, 6):
            cell = ws2.cell(row=5, column=col_idx)
            styles2[col_idx] = {
                'font': cell.font,
                'fill': cell.fill,
                'border': cell.border,
                'alignment': cell.alignment,
                'number_format': cell.number_format
            }

        # Shift existing data in columns A to F (columns 1 to 6) down by N rows
        max_r = ws2.max_row
        for r in range(max_r, 4, -1):
            for col_idx in range(1, 7):
                src_cell = ws2.cell(row=r, column=col_idx)
                dest_cell = ws2.cell(row=r + N, column=col_idx)
                dest_cell.value = src_cell.value
                if src_cell.has_style:
                    dest_cell.font = copy(src_cell.font)
                    dest_cell.fill = copy(src_cell.fill)
                    dest_cell.border = copy(src_cell.border)
                    dest_cell.alignment = copy(src_cell.alignment)
                    dest_cell.number_format = src_cell.number_format
                src_cell.value = None

        # Write data to Sheet2 (columns A to E)
        for idx, row in enumerate(new_rows):
            r_idx = 5 + idx
            dt = row['date']
            ws2.cell(row=r_idx, column=1, value=datetime.datetime.combine(dt, datetime.time.min))
            ws2.cell(row=r_idx, column=2, value=cum_values[dt]['ACB'])
            ws2.cell(row=r_idx, column=3, value=cum_values[dt]['VIB'])
            ws2.cell(row=r_idx, column=4, value=cum_values[dt]['VPB'])
            ws2.cell(row=r_idx, column=5, value=cum_values[dt]['TCB'])
            ws2.cell(row=r_idx, column=6, value='')

            # Apply style to columns 1 to 6 of the new rows
            for col_idx in range(1, 7):
                cell = ws2.cell(row=r_idx, column=col_idx)
                s = styles2.get(col_idx) or styles2.get(1)
                if s:
                    if s['font']: cell.font = copy(s['font'])
                    if s['fill']: cell.fill = copy(s['fill'])
                    if s['border']: cell.border = copy(s['border'])
                    if s['alignment']: cell.alignment = copy(s['alignment'])
                    if s['number_format']: cell.number_format = s['number_format']

    # Keep all existing labels and narrative text in the workbook untouched.
    # In particular, do not rebuild the month/YTD summary block in Sheet2 G:K.
    if N > 0:
        force_excel_recalculation(wb_dest)
        save_output_file(wb_dest, template_path)
        print(f"Successfully processed NN ban rong with {N} new rows; existing text was preserved.")
    else:
        print("No new dates for NN ban rong; workbook text was left unchanged.")
    return True

# ==================== GD TU DOANH & TIN DOANH NGHIEP ====================

import requests
import time

class FireantClient:
    BASE_URL = "https://api.fireant.vn"
    
    def __init__(self):
        self.session = requests.Session()
        self.token = None
        self.headers = {
            "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.114 Safari/537.36",
            "Content-Type": "application/json"
        }

    def login(self):
        url = f"{self.BASE_URL}/authentication/anonymous-login"
        try:
            resp = self.session.post(url, json={}, headers=self.headers, timeout=10)
            if resp.status_code == 200:
                data = resp.json()
                self.token = data.get("accessToken") or data.get("token")
                if self.token:
                    print("✅ Fireant Login Successful.")
                    self.headers["Authorization"] = f"Bearer {self.token}"
                else:
                    print("⚠️ Fireant Login Failed: No token found.")
            else:
                print(f"⚠️ Fireant Login Failed: {resp.status_code}")
        except Exception as e:
            print(f"⚠️ Fireant Login Error: {e}")

    def get_historical_quotes(self, symbol, start_date, end_date):
        if not self.token:
            self.login()
            
        url = f"{self.BASE_URL}/symbols/{symbol}/historical-quotes"
        params = {
            "startDate": start_date.strftime("%Y-%m-%d"),
            "endDate": end_date.strftime("%Y-%m-%d"),
            "offset": 0,
            "limit": 100
        }
        
        try:
            for _ in range(3):
                resp = self.session.get(url, params=params, headers=self.headers, timeout=10)
                if resp.status_code == 200:
                    return resp.json()
                elif resp.status_code == 401:
                    print("🔄 Token expired, refreshing...")
                    self.login()
                    time.sleep(1)
                elif resp.status_code == 429:
                    time.sleep(1)
                else:
                    break
            return None
        except Exception as e:
            print(f"   ⚠️ API Error for {symbol}: {e}")
            return None

class StockProcessor:
    def __init__(self, workspace_path):
        self.workspace_path = workspace_path
        self.pdf_folder = os.path.join(workspace_path, "gd tu doanh", "Du lieu tu doanh")
        self.output_path = find_latest_template(
            os.path.join(workspace_path, "gd tu doanh"), 
            "Thong ke Tu doanh Co phieu_update *.xlsx", 
            "Thong ke Tu doanh Co phieu.xlsx"
        )
        
        self.col_definitions = [
            {"name": "MB_Vol", "min": 170.0, "max": 190.0},
            {"name": "MS_Vol", "min": 240.0, "max": 265.0},
            {"name": "MB_Val", "min": 315.0, "max": 338.0},
            {"name": "MS_Val", "min": 390.0, "max": 412.0},
            {"name": "AB_Vol", "min": 460.0, "max": 485.0},
            {"name": "AS_Vol", "min": 535.0, "max": 558.0},
            {"name": "AB_Val", "min": 605.0, "max": 632.0},
            {"name": "AS_Val", "min": 680.0, "max": 705.0},
        ]

    def extract_from_pdf(self, pdf_path):
        all_data = []
        try:
            with pdfplumber.open(pdf_path) as pdf:
                p0_text = pdf.pages[0].extract_text() or ""
                date_match = re.search(r'Ngày\s+(\d{2}/\d{2}/\d{4})', p0_text)
                t_date = date_match.group(1) if date_match else "Unknown"
                
                if t_date == "Unknown":
                    basename = os.path.basename(pdf_path)
                    file_date_match = re.search(r'(\d{8})', basename)
                    if file_date_match:
                        d_str = file_date_match.group(1)
                        t_date = f"{d_str[6:8]}/{d_str[4:6]}/{d_str[0:4]}"

                stock_pages = []
                for idx, page in enumerate(pdf.pages):
                    text = page.extract_text() or ""
                    if "CHỨNG KHOÁN KHÁC" in text.upper() and "TỰ DOANH" in text.upper():
                        break
                    stock_pages.append(idx)
                
                for p_idx in stock_pages:
                    page = pdf.pages[p_idx]
                    words = page.extract_words()
                    
                    from collections import defaultdict
                    lines = defaultdict(list)
                    for w in words:
                        found = False
                        for y in lines:
                            if abs(w['top'] - y) < 4:
                                lines[y].append(w)
                                found = True
                                break
                        if not found:
                            lines[w['top']].append(w)
                    
                    for y in sorted(lines.keys()):
                        line_words = sorted(lines[y], key=lambda x: x['x0'])
                        
                        symbol = None
                        symbol_idx = -1
                        for idx, w in enumerate(line_words):
                            text = w['text'].strip()
                            if len(text) == 3 and text.isalpha() and text.isupper():
                                symbol = text
                                symbol_idx = idx
                                break
                        
                        if symbol and symbol_idx > 0:
                            prev_text = line_words[symbol_idx - 1]['text'].strip()
                            if prev_text.isdigit():
                                row_data = {"Date": t_date, "Symbol": symbol}
                                for col in self.col_definitions:
                                    row_data[col["name"]] = 0.0
                                    
                                for w in line_words[symbol_idx + 1:]:
                                    text = w['text'].replace(',', '').replace('.', '').strip()
                                    if not text:
                                        continue
                                    try:
                                        val = float(text)
                                    except ValueError:
                                        continue
                                    
                                    for col in self.col_definitions:
                                        if col["min"] <= w['x1'] <= col["max"]:
                                            row_data[col["name"]] = val
                                            break
                                            
                                all_data.append(row_data)
        except Exception as e:
            print(f"❌ Error processing PDF {pdf_path}: {e}")
            
        return pd.DataFrame(all_data)

    def run(self):
        print("🔍 Scanning PDFs in folder 'Du lieu tu doanh'...")
        if not os.path.exists(self.pdf_folder):
            print(f"❌ Folder not found: {self.pdf_folder}")
            return False

        pdf_files = sorted([f for f in os.listdir(self.pdf_folder) if f.endswith('.pdf')])
        print(f"📄 Found {len(pdf_files)} PDF files.")
        if not pdf_files:
            return False

        new_dfs = []
        for f in pdf_files:
            pdf_path = os.path.join(self.pdf_folder, f)
            print(f"   Parsing: {f}")
            df = self.extract_from_pdf(pdf_path)
            if not df.empty:
                new_dfs.append(df)

        if not new_dfs:
            print("⚠️ No data extracted from PDFs.")
            return False

        new_daily_df = pd.concat(new_dfs, ignore_index=True)
        print(f"✅ Extracted {len(new_daily_df)} daily trade records from PDFs.")

        existing_daily_df = pd.DataFrame()
        if os.path.exists(self.output_path):
            try:
                existing_daily_df = pd.read_excel(self.output_path, sheet_name="Daily_Data")
                print(f"📥 Loaded {len(existing_daily_df)} existing records from Excel.")
            except Exception as e:
                print(f"⚠️ Could not load existing Excel: {e}. Will create a new one.")

        if not existing_daily_df.empty:
            combined_daily_df = pd.concat([existing_daily_df, new_daily_df], ignore_index=True)
        else:
            combined_daily_df = new_daily_df

        original_len = len(combined_daily_df)
        combined_daily_df.drop_duplicates(subset=["Date", "Symbol"], keep="last", inplace=True)
        deduplicated_len = len(combined_daily_df)
        print(f"🧹 Deduplicated: {original_len} -> {deduplicated_len} daily records.")

        combined_daily_df['DateObj'] = pd.to_datetime(combined_daily_df['Date'], format="%d/%m/%Y", errors='coerce')
        combined_daily_df.sort_values(by=["DateObj", "Symbol"], ascending=[True, True], inplace=True)
        combined_daily_df.drop(columns=['DateObj'], inplace=True, errors='ignore')

        print("📊 Aggregating data by week...")
        temp_df = combined_daily_df.copy()
        temp_df['DateObj'] = pd.to_datetime(temp_df['Date'], format="%d/%m/%Y", errors='coerce')
        temp_df['Week_Start'] = temp_df['DateObj'] - temp_df['DateObj'].dt.weekday.map(lambda x: datetime.timedelta(days=x))
        
        def get_week_label(row):
            mon = row['Week_Start']
            fri = mon + datetime.timedelta(days=4)
            return f"{mon.strftime('%d/%m/%Y')} - {fri.strftime('%d/%m/%Y')}"
            
        temp_df['Week_Label'] = temp_df.apply(get_week_label, axis=1)
        sum_cols = ["MB_Vol", "MS_Vol", "MB_Val", "MS_Val", "AB_Vol", "AS_Vol", "AB_Val", "AS_Val"]
        weekly_grouped = temp_df.groupby(["Week_Start", "Week_Label", "Symbol"])[sum_cols].sum().reset_index()

        weekly_grouped["Total_Buy_Vol"] = weekly_grouped["MB_Vol"] + weekly_grouped["AB_Vol"]
        weekly_grouped["Total_Sell_Vol"] = weekly_grouped["MS_Vol"] + weekly_grouped["AS_Vol"]
        weekly_grouped["Total_Buy_Val"] = weekly_grouped["MB_Val"] + weekly_grouped["AB_Val"]
        weekly_grouped["Total_Sell_Val"] = weekly_grouped["MS_Val"] + weekly_grouped["AS_Val"]
        weekly_grouped["Net_Vol"] = weekly_grouped["Total_Buy_Vol"] - weekly_grouped["Total_Sell_Vol"]
        weekly_grouped["Net_Val"] = weekly_grouped["Total_Buy_Val"] - weekly_grouped["Total_Sell_Val"]

        weekly_grouped.sort_values(by=["Week_Start", "Symbol"], ascending=[False, True], inplace=True)

        latest_week_start = weekly_grouped["Week_Start"].max()
        latest_week_label = weekly_grouped[weekly_grouped["Week_Start"] == latest_week_start]["Week_Label"].iloc[0]
        print(f"📈 Processing Top 10 Tables for week: {latest_week_label}")

        latest_week_df = weekly_grouped[weekly_grouped["Week_Start"] == latest_week_start].copy()

        fireant = FireantClient()
        fireant.login()

        mon_date = latest_week_start
        fri_date = latest_week_start + datetime.timedelta(days=4)
        unique_symbols = latest_week_df["Symbol"].unique()
        valid_symbol_quotes = {}
        
        print("🔗 Querying Fireant API for weekly market data of all symbols...")
        for sym in unique_symbols:
            quotes = fireant.get_historical_quotes(sym, mon_date, fri_date)
            if quotes:
                valid_symbol_quotes[sym] = quotes

        buy_ratios = []
        sell_ratios = []
        close_prices = []

        for idx, row in weekly_grouped.iterrows():
            symbol = row["Symbol"]
            quotes = valid_symbol_quotes.get(symbol)
            w_start = row["Week_Start"]
            w_mon = w_start
            w_fri = w_start + datetime.timedelta(days=4)
            
            if w_start != latest_week_start:
                quotes = fireant.get_historical_quotes(symbol, w_mon, w_fri)

            weekly_total_val = 0.0
            last_close = 0.0
            
            if quotes:
                weekly_total_val = sum([q.get("totalValue", 0) for q in quotes if q.get("totalValue")])
                quotes_sorted = sorted(quotes, key=lambda q: q.get("date", ""))
                if quotes_sorted:
                    last_quote = quotes_sorted[-1]
                    last_close = (last_quote.get("priceClose") or last_quote.get("close", 0)) * 1000.0
            
            buy_val_vnd = row["Total_Buy_Val"] * 1000.0
            sell_val_vnd = row["Total_Sell_Val"] * 1000.0
            
            buy_ratio = buy_val_vnd / weekly_total_val if weekly_total_val > 0 else 0.0
            sell_ratio = sell_val_vnd / weekly_total_val if weekly_total_val > 0 else 0.0
            
            buy_ratios.append(buy_ratio)
            sell_ratios.append(sell_ratio)
            close_prices.append(last_close)

        weekly_grouped["Giá Đóng Cửa"] = close_prices
        weekly_grouped["Tỷ trọng Mua (%)"] = buy_ratios
        weekly_grouped["Tỷ trọng Bán (%)"] = sell_ratios

        validated_week_df = latest_week_df[latest_week_df["Symbol"].isin(valid_symbol_quotes.keys())].copy()
        top_sold = validated_week_df.sort_values(by="Net_Val", ascending=True).head(10).copy()
        top_bought = validated_week_df.sort_values(by="Net_Val", ascending=False).head(10).copy()

        def process_top_table(df_top, is_sell_table):
            processed_rows = []
            for idx, row in df_top.iterrows():
                symbol = row["Symbol"]
                quotes = valid_symbol_quotes.get(symbol, [])
                weekly_total_val = 0.0
                last_close = 0.0
                
                if quotes:
                    weekly_total_val = sum([q.get("totalValue", 0) for q in quotes if q.get("totalValue")])
                    quotes_sorted = sorted(quotes, key=lambda q: q.get("date", ""))
                    if quotes_sorted:
                        last_quote = quotes_sorted[-1]
                        last_close = (last_quote.get("priceClose") or last_quote.get("close", 0)) * 1000.0
                
                buy_val_vnd = row["Total_Buy_Val"] * 1000.0
                sell_val_vnd = row["Total_Sell_Val"] * 1000.0
                net_val_vnd = row["Net_Val"] * 1000.0

                ratio = 0.0
                if weekly_total_val > 0:
                    if is_sell_table:
                        ratio = sell_val_vnd / weekly_total_val
                    else:
                        ratio = buy_val_vnd / weekly_total_val

                processed_rows.append({
                    "Mã": symbol,
                    "Giá": last_close,
                    "Bán ròng (GT)" if is_sell_table else "Mua ròng (GT)": net_val_vnd,
                    "Tỷ trọng bán của Khối Tự doanh" if is_sell_table else "Tỷ trọng mua của Khối Tự doanh": ratio
                })
            return pd.DataFrame(processed_rows)

        top_sold_processed = process_top_table(top_sold, is_sell_table=True)
        top_bought_processed = process_top_table(top_bought, is_sell_table=False)

        weekly_grouped_renamed = weekly_grouped.copy()
        weekly_grouped_renamed.rename(columns={
            "Week_Label": "Tuần",
            "Symbol": "Mã CK",
            "MB_Vol": "KL Mua Khớp Lệnh",
            "MS_Vol": "KL Bán Khớp Lệnh",
            "MB_Val": "GT Mua Khớp Lệnh",
            "MS_Val": "GT Bán Khớp Lệnh",
            "AB_Vol": "KL Mua Thỏa Thuận",
            "AS_Vol": "KL Bán Thỏa Thuận",
            "AB_Val": "GT Mua Thỏa Thuận",
            "AS_Val": "GT Bán Thỏa Thuận",
            "Total_Buy_Vol": "Tổng KL Mua",
            "Total_Sell_Vol": "Tổng KL Bán",
            "Total_Buy_Val": "Tổng GT Mua",
            "Total_Sell_Val": "Tổng GT Bán",
            "Net_Vol": "KL Ròng",
            "Net_Val": "GT Ròng"
        }, inplace=True)
        weekly_grouped_renamed.drop(columns=["Week_Start"], inplace=True)

        clean_template_path = os.path.join(self.workspace_path, "gd tu doanh", "Thong ke Tu doanh Co phieu.xlsx")
        
        dir_name = os.path.dirname(clean_template_path)
        base_name = os.path.basename(clean_template_path)
        name, ext = os.path.splitext(base_name)
        suffix = get_update_suffix()
        output_path = os.path.join(dir_name, f"{name}{suffix}{ext}")

        print(f"💾 Saving to {output_path}...")
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            combined_daily_df.to_excel(writer, sheet_name="Daily_Data", index=False)
            weekly_grouped_renamed.to_excel(writer, sheet_name="Weekly_Summary", index=False)
            top_sold_processed.to_excel(writer, sheet_name="Top_10_Ban_Rong", index=False)
            top_bought_processed.to_excel(writer, sheet_name="Top_10_Mua_Rong", index=False)

        print(f"Saved processed output to: {output_path}")
        
        # Delete clean template file if it exists and is different from the output path
        if os.path.exists(clean_template_path) and os.path.abspath(clean_template_path) != os.path.abspath(output_path):
            try:
                os.remove(clean_template_path)
                print(f"Deleted clean template file: {os.path.basename(clean_template_path)}")
            except OSError as e:
                print(f"Warning: Could not delete clean template file {clean_template_path}: {e}")
                
        # Delete old update files matching the base pattern in this directory
        old_updates = glob.glob(os.path.join(dir_name, f"{name}_update *.xlsx"))
        for p in old_updates:
            if os.path.abspath(p) != os.path.abspath(output_path):
                try:
                    os.remove(p)
                    print(f"Deleted old update file: {os.path.basename(p)}")
                except OSError as e:
                    print(f"Warning: Could not delete old update file {p}: {e}")
                    
        print("🎉 Processing completed successfully!")
        return True

def run_tu_doanh_feature(workspace_dir):
    print("\n================ RUNNING GD TU DOANH FEATURE ================")
    # The weekly FiinProX exports already contain the final Top 10 values.
    # Do not rebuild them from ROOM PDFs or query Fireant: doing so duplicates
    # work, changes the source definition, and fails when a PDF date is missing.
    try:
        import json
        from weekly_report import ensure_proprietary_workbook, latest_file

        config_path = os.path.join(workspace_dir, "report_config.json")
        with open(config_path, "r", encoding="utf-8") as config_file:
            config = json.load(config_file)

        patterns = config.get("proprietary_raw_patterns", {})
        buy_path = latest_file(patterns.get("buy", []))
        sell_path = latest_file(patterns.get("sell", []))
        if not buy_path or not sell_path:
            print("ERROR: Missing proprietary-trading Top 10 Buy or Top 10 Sell Excel input.")
            print("Place both FiinProX domestic-institution Top gia tri rong files in incoming and run again.")
            return False

        source_dates = []
        for source_path in (buy_path, sell_path):
            match = re.search(r"(20\d{6})(?=\.xlsx$)", source_path.name, re.IGNORECASE)
            if match:
                source_dates.append(datetime.datetime.strptime(match.group(1), "%Y%m%d").date())
        as_of = max(source_dates) if source_dates else get_report_as_of()

        output_path = ensure_proprietary_workbook(as_of, config)
        if not output_path:
            print("ERROR: Could not create the proprietary-trading Top 10 workbook.")
            return False

        print(f"Copied Top 10 Sell from: {sell_path.name}")
        print(f"Copied Top 10 Buy from:  {buy_path.name}")
        print(f"Saved proprietary-trading tables to: {output_path}")
        return True
    except Exception as exc:
        print(f"ERROR: Could not build proprietary-trading Top 10 tables: {exc}")
        return False

def run_tin_doanh_nghiep_feature(workspace_dir, watchlist=None):
    print("\n================ RUNNING TIN DOANH NGHIEP FEATURE ================")
    if watchlist is None:
        watchlist = ['ACB', 'TCB', 'VIB', 'VPB']
        
    print(f"Watchlist: {', '.join(watchlist)}")
    
    reference_date = datetime.datetime.combine(get_report_as_of(), datetime.time.min)
    current_weekday = reference_date.weekday()
    days_to_friday = 4 - current_weekday
    this_friday = reference_date + datetime.timedelta(days=days_to_friday)
    end_date = this_friday.replace(hour=17, minute=0, second=0, microsecond=0)
    last_saturday = this_friday - datetime.timedelta(days=6)
    start_date = last_saturday.replace(hour=0, minute=0, second=0, microsecond=0)
    
    print(f"Khung thời gian: {start_date.strftime('%Y-%m-%d %H:%M:%S')} đến {end_date.strftime('%Y-%m-%d %H:%M:%S')}")
    
    print(f"[*] Thu thập Tin chính thống từ {start_date.strftime('%Y-%m-%d')} đến {end_date.strftime('%Y-%m-%d')}...")
    raw_news = []
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
    }
    
    keywords = ["ĐHĐCĐ", "sở hữu", "cổ tức", "giao dịch cổ phiếu", "phát hành"]
    
    for ticker in watchlist:
        url = f"https://finfo-api.ipas.com.vn/v4/news?q=tagCodes:{ticker}~newsSource:HOSE,HNX,UPCOM&sort=newsDate:desc~newsTime:desc&size=50&page=1"
        try:
            response = requests.get(url, headers=headers, timeout=10)
            response.raise_for_status()
            data = response.json().get('data', [])
            
            for item in data:
                title = item.get('newsTitle', '')
                link = item.get('newsUrl', '')
                time_str = item.get('newsDate', '')
                attachments = item.get('attachments', [])
                pdf_url = ""
                if attachments and len(attachments) > 0:
                    pdf_url = attachments[0].get('url', '')
                
                title_lower = title.lower()
                matched_keyword = ""
                for kw in keywords:
                    if kw.lower() in title_lower:
                        matched_keyword = kw
                        break
                        
                if time_str:
                    try:
                        if len(time_str) >= 10:
                            pub_date = datetime.datetime.strptime(time_str[:10], "%Y-%m-%d")
                        else:
                            pub_date = datetime.datetime.combine(get_report_as_of(), datetime.time.min)
                            
                        if start_date.date() <= pub_date.date() <= end_date.date():
                            raw_news.append({
                                "Mã CK": ticker,
                                "Thời gian": time_str,
                                "Loại tin": "Chính thống (Sở Giao Dịch)",
                                "Phân loại": matched_keyword,
                                "Tiêu đề gốc": title,
                                "Nội dung tóm tắt": "Chưa tóm tắt (AI đang tắt)",
                                "Tác động": "Trung tính",
                                "Link nguồn": link,
                                "PDF Link": pdf_url
                            })
                    except Exception as e:
                        print(f"[!] Lỗi parse ngày tháng {time_str}: {e}")
        except Exception as e:
            print(f"[!] Lỗi khi lấy tin tức cho mã {ticker} qua API: {e}")
        time.sleep(1)
        
    if not raw_news:
        print("⚠️ Không thu thập được tin tức chính thống nào.")
        return False
        
    df = pd.DataFrame(raw_news)
    df.drop_duplicates(subset=['Tiêu đề gốc'], keep='first', inplace=True)
    cleaned_news = df.to_dict('records')
    
    processed_news = []
    filtered_news = []
    for news in cleaned_news:
        processed_news.append(news)
        if news.get("Phân loại"):
            filtered_news.append(news)
            
    output_dir = os.path.join(workspace_dir, "tin doanh nghiep")
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
        
    tickers_str = "_".join(watchlist)
    start_str = start_date.strftime("%d")
    end_str = end_date.strftime("%d_%b_%Y")
    
    suffix = get_update_suffix()
    filename = f"BaoCao_TinChinhThong_{tickers_str}_Tuan_{start_str}_{end_str}{suffix}.xlsx"
    filepath = os.path.join(output_dir, filename)
    
    print(f"[*] Đang xuất báo cáo ra file Excel: {filepath}...")
    
    df_dict = {
        'Tin_Chinh_Thong_Toan_Bo': pd.DataFrame(processed_news),
        'Tin_Chinh_Thong_Loc': pd.DataFrame(filtered_news)
    }
    
    from openpyxl.styles import Font
    with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
        for sheet_name, sheet_df in df_dict.items():
            if not sheet_df.empty:
                sheet_df.to_excel(writer, sheet_name=sheet_name, index=False)
            else:
                pd.DataFrame(columns=["Không có dữ liệu"]).to_excel(writer, sheet_name=sheet_name, index=False)
        
        workbook = writer.book
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            for cell in worksheet[1]:
                cell.font = Font(name=REPORT_FONT, bold=True)
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name=REPORT_FONT, size=cell.font.sz or 11)
            for column_cells in worksheet.columns:
                max_length = 0
                column = column_cells[0].column_letter
                for cell in column_cells:
                    try:
                        if cell.value and len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                adjusted_width = (max_length + 2)
                worksheet.column_dimensions[column].width = min(adjusted_width, 80)
                
    print(f"Saved processed output to: {filepath}")
    
    # Delete old report files in this directory
    old_reports = glob.glob(os.path.join(output_dir, "BaoCao_TinChinhThong_*_update *.xlsx"))
    for p in old_reports:
        if os.path.abspath(p) != os.path.abspath(filepath):
            try:
                os.remove(p)
                print(f"Deleted old report file: {os.path.basename(p)}")
            except OSError as e:
                print(f"Warning: Could not delete old report file {p}: {e}")
                
    print("✅ HOÀN THÀNH QUÁ TRÌNH TẠO BÁO CÁO TIN DOANH NGHIỆP!\n")
    return True

CALENDAR_COUNTRY_TRANSLATIONS = {
    "Australia": "Úc",
    "Canada": "Canada",
    "China": "Trung Quốc",
    "Euro Area": "Khu vực Euro",
    "France": "Pháp",
    "Germany": "Đức",
    "Global": "Toàn cầu",
    "India": "Ấn Độ",
    "Italy": "Ý",
    "Japan": "Nhật Bản",
    "New Zealand": "New Zealand",
    "South Korea": "Hàn Quốc",
    "United Kingdom": "Anh",
    "United States": "Mỹ",
}


def translate_calendar_event_name(event_name):
    """Translate the small, filtered macro-event vocabulary deterministically."""
    text = re.sub(r'\s+', ' ', str(event_name or '')).strip()
    if not text:
        return text

    lower = text.casefold()
    period = ""
    for token, translated in (
        ("mom", "theo tháng (MoM)"),
        ("yoy", "theo năm (YoY)"),
        ("qoq", "theo quý (QoQ)"),
    ):
        if re.search(rf'\b{token}\b', lower, flags=re.IGNORECASE):
            period = translated
            text = re.sub(rf'\b{token}\b', '', text, flags=re.IGNORECASE).strip()
            lower = text.casefold()
            break

    qualifier = ""
    qualifier_rules = (
        (r'\b(prel(?:iminary)?)\b', "sơ bộ"),
        (r'\bflash\b', "ước tính nhanh"),
        (r'\bfinal\b', "chính thức"),
    )
    for pattern, translated in qualifier_rules:
        if re.search(pattern, text, flags=re.IGNORECASE):
            qualifier = translated
            text = re.sub(pattern, '', text, flags=re.IGNORECASE).strip()
            break

    lower = text.casefold()
    if "non-farm" in lower or "non farm" in lower or "nonfarm" in lower:
        translated_name = "Bảng lương phi nông nghiệp"
    elif "core inflation rate" in lower:
        translated_name = "Lạm phát cơ bản"
    elif "inflation rate" in lower:
        translated_name = "Lạm phát"
    elif "gdp growth" in lower:
        translated_name = "Tăng trưởng GDP"
    elif "interest rate" in lower:
        authority = re.sub(
            r'\binterest rate(?: decision)?\b',
            '',
            text,
            flags=re.IGNORECASE,
        ).strip(' -')
        translated_name = f"Quyết định lãi suất của {authority}" if authority else "Quyết định lãi suất"
    else:
        # The feature filters to the vocabulary above.  Keeping an unexpected
        # source label is safer than inventing a translation for an unknown term.
        translated_name = text

    suffixes = [value for value in (period, qualifier) if value]
    return f"{translated_name}, {', '.join(suffixes)}" if suffixes else translated_name


def group_calendar_events_by_date(events):
    grouped = {}
    for event in events:
        event_date = event["Ngày"]
        grouped.setdefault(event_date, [])
        if event["Dữ liệu"] not in grouped[event_date]:
            grouped[event_date].append(event["Dữ liệu"])
    return [
        {"Ngày": event_date, "Dữ liệu": "\n".join(f"• {item}" for item in grouped[event_date])}
        for event_date in sorted(grouped)
    ]


def run_lich_su_kien_feature(workspace_dir):
    print("\n================ RUNNING LICH SU KIEN FEATURE ================")
    output_dir = os.path.join(workspace_dir, "lich su kien")
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
        print(f"Created output directory: {output_dir}")

    # 1. Calculate next week's Monday and Sunday
    today = get_report_as_of()
    days_to_next_monday = 7 - today.weekday()
    next_monday = today + datetime.timedelta(days=days_to_next_monday)
    next_sunday = next_monday + datetime.timedelta(days=6)

    start_str = next_monday.strftime('%Y-%m-%d')
    end_str = next_sunday.strftime('%Y-%m-%d')
    print(f"Lọc sự kiện tuần tới từ {next_monday.strftime('%d/%m/%Y')} đến {next_sunday.strftime('%d/%m/%Y')}...")

    # 2. Fetch the calendar page
    headers = {
        "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,image/apng,*/*;q=0.8",
        "Accept-Language": "en-US,en;q=0.9",
    }
    url = f"https://tradingeconomics.com/calendar?start={start_str}&end={end_str}"
    print(f"Fetching calendar data from: {url}...")
    try:
        from bs4 import BeautifulSoup
        resp = requests.get(url, headers=headers, timeout=15)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, 'html.parser')
    except Exception as e:
        print(f"❌ Error fetching calendar from Trading Economics: {e}")
        return False

    calendar_table = soup.find('table', id='calendar')
    if not calendar_table:
        print("❌ Could not find calendar table on Trading Economics.")
        return False

    rows = calendar_table.find_all('tr')
    
    keywords = [
        "inflation rate",
        "interest rate decision",
        "gdp growth",
        "non-farm employment"
    ]

    events_extracted = []

    for row in rows:
        if row.has_attr('data-id'):
            cells = row.find_all('td')
            if not cells:
                continue
                
            # Get event date from the class of the first cell
            date_class = cells[0].get('class', [])
            if not date_class:
                continue
            event_date_str = date_class[0]  # e.g. "2026-06-19"
            
            try:
                event_date_parsed = datetime.datetime.strptime(event_date_str, "%Y-%m-%d").date()
            except ValueError:
                continue
                
            # Filter 1: Next week only
            if not (next_monday <= event_date_parsed <= next_sunday):
                continue
                
            # Filter 2: High impact (class calendar-date-3)
            span_high = cells[0].find('span', class_='calendar-date-3')
            if not span_high:
                continue
                
            # Get event name
            event_link = row.find('a', class_='calendar-event')
            if not event_link:
                continue
            event_name = event_link.get_text(strip=True)
            
            # Get country
            country = row.get('data-country', '').strip().title()
            if not country:
                # Fallback to flag title
                flag_div = row.find('div', class_=lambda x: x and 'flag-' in x)
                if flag_div:
                    country = flag_div.get('title', '').strip().title()
            if not country:
                country = "Global"

            # Filter 3: Keywords (case-insensitive substring)
            event_name_lower = event_name.lower()
            matched = False
            for kw in keywords:
                if kw == "non-farm employment":
                    if any(x in event_name_lower for x in ["non-farm employment", "non farm employment", "nonfarm employment", "non-farm payroll", "non farm payroll", "nonfarm payroll"]):
                        matched = True
                        break
                elif kw == "interest rate decision":
                    if any(x in event_name_lower for x in ["interest rate decision", "interest rate"]):
                        matched = True
                        break
                elif kw in event_name_lower:
                    matched = True
                    break
                    
            if matched:
                country_vi = CALENDAR_COUNTRY_TRANSLATIONS.get(country, country)
                events_extracted.append({
                    "Ngày": datetime.datetime.combine(event_date_parsed, datetime.time.min),
                    "Dữ liệu": f"[{country_vi}] {translate_calendar_event_name(event_name)}"
                })

    print(f"Extracted {len(events_extracted)} high-impact important events for next week.")
    events_extracted = group_calendar_events_by_date(events_extracted)
    print(f"Combined events into {len(events_extracted)} date rows.")

    # 3. Create Excel workbook and save
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lịch sự kiện"

    # Set headers
    ws['A1'] = "Ngày"
    ws['B1'] = "Dữ liệu"

    # Styling headers
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF')
    )

    ws['A1'].fill = header_fill
    ws['A1'].font = header_font
    ws['A1'].alignment = center_align
    ws['A1'].border = thin_border

    ws['B1'].fill = header_fill
    ws['B1'].font = header_font
    ws['B1'].alignment = left_align
    ws['B1'].border = thin_border

    # Write data rows
    data_font = Font(name=REPORT_FONT, size=11)
    
    for r_idx, ev in enumerate(events_extracted, start=2):
        cell_date = ws.cell(row=r_idx, column=1, value=ev["Ngày"])
        cell_date.number_format = 'dd/mm/yyyy'
        cell_date.font = data_font
        cell_date.alignment = center_align
        cell_date.border = thin_border

        cell_data = ws.cell(row=r_idx, column=2, value=ev["Dữ liệu"])
        cell_data.font = data_font
        cell_data.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        cell_data.border = thin_border
        ws.row_dimensions[r_idx].height = max(18, 15 * (str(ev["Dữ liệu"]).count('\n') + 1))

    # Adjust column widths
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 60

    output_path = os.path.join(output_dir, "Lich su kien.xlsx")
    save_output_file(wb, output_path)
    print("✅ HOÀN THÀNH QUÁ TRÌNH TẠO BÁO CÁO LỊCH SỰ KIỆN!\n")
    return True

# ==================== MAIN ====================

MODULES = [
    ("2", "Ngành", run_nganh_feature),
    ("3", "Danh mục MBS", run_mbs_feature),
    ("4", "Giao dịch Nước ngoài", run_foreign_feature),
    ("5", "Thanh khoản thị trường", run_thanh_khoan_feature),
    ("6", "Đóng góp của Vingroup", run_vingroup_feature),
    ("7", "Nước ngoài bán ròng", run_nn_ban_rong_feature),
    ("8", "Giao dịch Tự doanh", run_tu_doanh_feature),
    ("9", "Tin doanh nghiệp", run_tin_doanh_nghiep_feature),
    ("10", "Lịch sự kiện", run_lich_su_kien_feature),
]


def run_selected_modules(workspace_dir, excluded=None):
    excluded = excluded or set()
    results = []
    for number, name, function in MODULES:
        if number in excluded:
            results.append((name, "BỎ QUA", "được loại trừ"))
            print(f"[BỎ QUA] {name}: được loại trừ")
            continue
        try:
            outcome = function(workspace_dir)
            status = "THÀNH CÔNG" if outcome is not False else "KHÔNG ĐẠT"
            detail = "hoàn tất" if outcome is not False else "thiếu dữ liệu hoặc điều kiện đầu vào"
        except Exception as exc:
            status, detail = "LỖI", str(exc)
        results.append((name, status, detail))
        print(f"[{status}] {name}: {detail}")
    failed = [name for name, status, _ in results if status in {"KHÔNG ĐẠT", "LỖI"}]
    if failed:
        print("\n[TỔNG HỢP] CÓ MODULE KHÔNG ĐẠT: " + ", ".join(failed))
        return False
    print("\n[TỔNG HỢP] TẤT CẢ MODULE ĐƯỢC CHỌN ĐỀU CHẠY THÀNH CÔNG")
    return True

def interactive_menu(workspace_dir):
    while True:
        print("\n================ CHƯƠNG TRÌNH XỬ LÝ DỮ LIỆU BCTT TUẦN ================")
        print("1. Chạy tất cả tính năng")
        print("2. Tính năng 'Ngành'")
        print("3. Tính năng 'Danh mục MBS'")
        print("4. Tính năng 'Giao dịch Nước ngoài' (GD Nuoc Ngoai)")
        print("5. Tính năng 'Thanh khoản thị trường' (Thanh khoan tt)")
        print("6. Tính năng 'Đóng góp của Vingroup'")
        print("7. Tính năng 'Nước ngoài bán ròng' (NN ban rong)")
        print("8. Tính năng 'Giao dịch Tự doanh' (GD tu doanh)")
        print("9. Tính năng 'Tin doanh nghiệp' (Tin doanh nghiep)")
        print("10. Tính năng 'Lịch sự kiện' (Lich su kien)")
        print("0. Thoát")
        print("======================================================================")
        try:
            choice = input("Nhập lựa chọn của bạn (0-10): ").strip()
        except (KeyboardInterrupt, EOFError):
            print("\nThoát chương trình.")
            break

        if choice == '0':
            print("Thoát chương trình.")
            break
        elif choice == '1':
            exclude_input = input("Nhập số thứ tự các tính năng muốn loại trừ (cách nhau bởi dấu phẩy, VD: 2,5), hoặc nhấn Enter để chạy tất cả: ").strip()
            excluded = set()
            if exclude_input:
                excluded = {x.strip() for x in exclude_input.split(',')}
            
            run_any = False
            if '2' not in excluded:
                run_nganh_feature(workspace_dir)
                run_any = True
            if '3' not in excluded:
                run_mbs_feature(workspace_dir)
                run_any = True
            if '4' not in excluded:
                run_foreign_feature(workspace_dir)
                run_any = True
            if '5' not in excluded:
                run_thanh_khoan_feature(workspace_dir)
                run_any = True
            if '6' not in excluded:
                run_vingroup_feature(workspace_dir)
                run_any = True
            if '7' not in excluded:
                run_nn_ban_rong_feature(workspace_dir)
                run_any = True
            if '8' not in excluded:
                run_tu_doanh_feature(workspace_dir)
                run_any = True
            if '9' not in excluded:
                run_tin_doanh_nghiep_feature(workspace_dir)
                run_any = True
            if '10' not in excluded:
                run_lich_su_kien_feature(workspace_dir)
                run_any = True
            
            if run_any:
                print("\nĐã xử lý xong các tính năng được chọn.")
            else:
                print("\nKhông có tính năng nào được chạy.")
            break
        elif choice == '2':
            run_nganh_feature(workspace_dir)
            print("\nĐã xử lý xong tính năng 'Ngành'.")
        elif choice == '3':
            run_mbs_feature(workspace_dir)
            print("\nĐã xử lý xong tính năng 'Danh mục MBS'.")
        elif choice == '4':
            run_foreign_feature(workspace_dir)
            print("\nĐã xử lý xong tính năng 'Giao dịch Nước ngoài'.")
        elif choice == '5':
            run_thanh_khoan_feature(workspace_dir)
            print("\nĐã xử lý xong tính năng 'Thanh khoản thị trường'.")
        elif choice == '6':
            run_vingroup_feature(workspace_dir)
            print("\nĐã xử lý xong tính năng 'Đóng góp của Vingroup'.")
        elif choice == '7':
            run_nn_ban_rong_feature(workspace_dir)
            print("\nĐã xử lý xong tính năng 'Nước ngoài bán ròng'.")
        elif choice == '8':
            run_tu_doanh_feature(workspace_dir)
            print("\nĐã xử lý xong tính năng 'Giao dịch Tự doanh'.")
        elif choice == '9':
            tickers_input = input("Nhập danh sách mã cổ phiếu cách nhau bằng dấu phẩy (VD: ACB,TCB,VIB) [Nhấn Enter để dùng mặc định ACB,TCB,VIB,VPB]: ").strip()
            if tickers_input:
                watchlist = [t.strip().upper() for t in tickers_input.split(',')]
            else:
                watchlist = ['ACB', 'TCB', 'VIB', 'VPB']
            run_tin_doanh_nghiep_feature(workspace_dir, watchlist)
            print("\nĐã xử lý xong tính năng 'Tin doanh nghiệp'.")
        elif choice == '10':
            run_lich_su_kien_feature(workspace_dir)
            print("\nĐã xử lý xong tính năng 'Lịch sự kiện'.")
        else:
            print("Lựa chọn không hợp lệ. Vui lòng nhập từ 0 đến 10.")

def main():
    workspace_dir = os.path.dirname(os.path.abspath(__file__))
    print(f"Workspace Directory: {workspace_dir}")
    
    if len(sys.argv) > 1:
        arg = sys.argv[1].lower()
        if arg == 'all':
            excluded = set()
            if len(sys.argv) > 2:
                excluded = {x.strip() for x in sys.argv[2].split(',')}
            elif '--exclude' in sys.argv:
                try:
                    idx = sys.argv.index('--exclude')
                    excluded = {x.strip() for x in sys.argv[idx+1].split(',')}
                except (ValueError, IndexError):
                    pass

            if not run_selected_modules(workspace_dir, excluded):
                raise SystemExit(1)
        elif 'nganh' in arg:
            if run_nganh_feature(workspace_dir) is False:
                raise SystemExit(1)
        elif 'mbs' in arg:
            if run_mbs_feature(workspace_dir) is False:
                raise SystemExit(1)
        elif 'foreign' in arg or 'nuocngoai' in arg:
            if run_foreign_feature(workspace_dir) is False:
                raise SystemExit(1)
        elif 'thanhkhoan' in arg:
            if run_thanh_khoan_feature(workspace_dir) is False:
                raise SystemExit(1)
        elif 'vingroup' in arg:
            if run_vingroup_feature(workspace_dir) is False:
                raise SystemExit(1)
        elif 'nnbanrong' in arg:
            if run_nn_ban_rong_feature(workspace_dir) is False:
                raise SystemExit(1)
        elif 'tudoanh' in arg:
            if run_tu_doanh_feature(workspace_dir) is False:
                raise SystemExit(1)
        elif 'tindoanhnghiep' in arg:
            if run_tin_doanh_nghiep_feature(workspace_dir) is False:
                raise SystemExit(1)
        elif 'lichsukien' in arg:
            if run_lich_su_kien_feature(workspace_dir) is False:
                raise SystemExit(1)
        else:
            print(f"Tham số không hợp lệ: {sys.argv[1]}. Đang mở menu lựa chọn...")
            interactive_menu(workspace_dir)
    else:
        interactive_menu(workspace_dir)

if __name__ == '__main__':
    main()
